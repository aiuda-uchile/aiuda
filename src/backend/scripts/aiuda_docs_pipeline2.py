#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import html
import io
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import multiprocessing

# Detectar dispositivo óptimo: GPU si está disponible, CPU con todos los cores si no
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")
_CPU_CORES = multiprocessing.cpu_count()
os.environ.setdefault("OMP_NUM_THREADS", str(_CPU_CORES))

APP_NAME = "AIUDA Docs Pipeline"
DEFAULT_TRANSLATION_MODEL = "facebook/nllb-200-distilled-600M"
ALLOWED_TARGETS = {"es", "en", "pt", "gl"}
NLLB_LANG_MAP = {
    "es": "spa_Latn",
    "en": "eng_Latn",
    "pt": "por_Latn",
    "gl": "glg_Latn",
}
SKIP_TRANSLATION_KEYS = {
    "task_id",
    "input_file",
    "input_filename",
    "document_type",
    "source_language_requested",
    "source_language_detected",
    "source_language_used",
    "report_language",
    "target_language",
    "filename",
    "unit_index",
    "id",
    "category_code",
    "severity_code",
    "excerpt",
    "excerpt_original",
    "issue_count",
    "page_count",
    "slide_count",
    "units_analyzed",
    "metric_key",
    "reference_code",
    "metric_type_code",
    "ratio",
    "score",
    "score_ratio",
    "threshold",
    "bbox",
    "figure_index",
    "text_block_index",
    "color_fg",
    "color_bg",
    "base_url",
    "issue_id",
    "preview_image_path",
    "preview_source_page",
}

WCAG_NORMAL_TEXT_RATIO = 4.5
WCAG_LARGE_TEXT_RATIO = 3.0
WCAG_NON_TEXT_RATIO = 3.0
TEXT_DENSITY_WORD_THRESHOLD = 120
PARAGRAPH_DENSITY_WORD_THRESHOLD = 55
SMALL_FONT_THRESHOLD_PT = 18.0
MAX_RENDER_SIDE = 1400

METHODOLOGY_REFERENCES = [
    {
        "reference_code": "WCAG-1.4.3",
        "title": "W3C WCAG 2.1 - Success Criterion 1.4.3 Contrast (Minimum)",
        "metric_type": "Umbral normativo de referencia",
        "details": "Se usa como referencia 4.5:1 para texto normal y 3:1 para texto grande.",
    },
    {
        "reference_code": "WCAG-1.4.11",
        "title": "W3C WCAG 2.1 - Success Criterion 1.4.11 Non-text Contrast",
        "metric_type": "Umbral normativo de referencia",
        "details": "Se usa como referencia 3:1 para componentes visuales y gráficos relevantes.",
    },
    {
        "reference_code": "WCAG2ICT",
        "title": "W3C WCAG2ICT - Applying WCAG to Non-Web Documents and Software",
        "metric_type": "Umbral normativo de referencia",
        "details": "Sirve de base para aplicar estos criterios a PDF, PPT y PPTX, aunque su naturaleza es orientativa.",
    },
    {
        "reference_code": "BRETTEL-1997",
        "title": "Brettel et al., 1997 - Computerized simulation of color appearance for dichromats",
        "metric_type": "Métrica heurística del sistema",
        "details": "Referencia metodológica para simulación/estimación de protanopia, deuteranopia y tritanopia.",
    },
    {
        "reference_code": "MACHADO-2009",
        "title": "Machado et al., 2009 - A Physiologically-based Model for Simulation of Color Vision Deficiency",
        "metric_type": "Métrica heurística del sistema",
        "details": "Referencia metodológica para simulación de deficiencia visual cromática y análisis comparativo orientativo.",
    },
]

CVD_MATRICES = {
    "protanopia": (
        (0.152286, 1.052583, -0.204868),
        (0.114503, 0.786281, 0.099216),
        (-0.003882, -0.048116, 1.051998),
    ),
    "deuteranopia": (
        (0.367322, 0.860646, -0.227968),
        (0.280085, 0.672501, 0.047413),
        (-0.011820, 0.042940, 0.968881),
    ),
    "tritanopia": (
        (1.255528, -0.076749, -0.178779),
        (-0.078411, 0.930809, 0.147602),
        (0.004733, 0.691367, 0.303900),
    ),
}


# ---------------------------------------------------------
# ARGUMENTS
# ---------------------------------------------------------


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="AIUDA - análisis documental y generación de informes multilingües"
    )
    parser.add_argument("input_file", type=str, help="Archivo PDF/PPT/PPTX de entrada")
    parser.add_argument(
        "-o",
        "--output-dir",
        type=str,
        required=True,
        help="Directorio de salida",
    )
    parser.add_argument(
        "--task-id",
        type=str,
        default="",
        help="ID de tarea del backend",
    )
    parser.add_argument(
        "--source-lang",
        type=str,
        default=None,
        help="Idioma del documento (es, en, pt, gl). Si no se indica, se intentará detectar.",
    )
    parser.add_argument(
        "--translate-to",
        nargs="*",
        default=["es", "en", "pt", "gl"],
        help="Idiomas destino para el informe. Ejemplo: --translate-to es en pt gl",
    )
    parser.add_argument(
        "--ocr",
        action="store_true",
        help="Intenta OCR cuando un PDF no contiene texto extraíble (solo se usa si hay soporte disponible).",
    )
    parser.add_argument(
        "--max-translation-chars",
        type=int,
        default=800,
        help="Tamaño máximo por bloque de traducción",
    )
    return parser.parse_args()


# ---------------------------------------------------------
# HELPERS
# ---------------------------------------------------------


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)



def normalize_whitespace(text: str) -> str:
    text = str(text or "")
    text = re.sub(r"[\u200b-\u200f\u202a-\u202e\u2060\ufeff\ufffe\uffff]", "", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def sanitize_stem(name: str) -> str:
    return re.sub(r"[^\w.\-]+", "_", name)



def normalize_lang_code(code: Optional[str]) -> Optional[str]:
    if not code:
        return None
    code = code.strip().lower()
    aliases = {
        "es-es": "es",
        "es_419": "es",
        "en-us": "en",
        "en-gb": "en",
        "pt-br": "pt",
        "pt-pt": "pt",
        "gl-es": "gl",
    }
    return aliases.get(code, code)



def write_text(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")



def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")



def collect_output_filenames(output_dir: Path) -> List[str]:
    return sorted([p.name for p in output_dir.iterdir() if p.is_file()])



def build_logger(output_dir: Path) -> logging.Logger:
    logger = logging.getLogger("aiuda_docs")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")

    sh = logging.StreamHandler(sys.stdout)
    sh.setLevel(logging.INFO)
    sh.setFormatter(formatter)

    fh = logging.FileHandler(output_dir / "job.log", encoding="utf-8")
    fh.setLevel(logging.INFO)
    fh.setFormatter(formatter)

    logger.addHandler(sh)
    logger.addHandler(fh)
    return logger



def write_status(
    output_dir: Path,
    task_id: str,
    state: str,
    progress: int,
    stage: str,
    message: str,
    extra: Optional[dict] = None,
) -> None:
    payload = {
        "task_id": task_id,
        "state": state,
        "progress": progress,
        "current_stage": stage,
        "message": message,
        "outputs": collect_output_filenames(output_dir),
    }
    if extra:
        payload.update(extra)
    write_json(output_dir / "status.json", payload)



def split_text_for_translation(text: str, max_chars: int) -> List[str]:
    text = normalize_whitespace(text)
    if not text:
        return []

    parts = re.split(r"(?<=[\.\!\?\:\;])\s+", text)
    parts = [p.strip() for p in parts if p.strip()]
    if not parts:
        parts = [text]

    chunks: List[str] = []
    current = ""

    for part in parts:
        candidate = part if not current else current + " " + part
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                chunks.append(current)
            if len(part) <= max_chars:
                current = part
            else:
                start = 0
                while start < len(part):
                    chunks.append(part[start:start + max_chars])
                    start += max_chars
                current = ""

    if current:
        chunks.append(current)

    return chunks



def detect_language_heuristic(text: str) -> str:
    text = normalize_whitespace(text).lower()
    if not text:
        return "es"

    sample = text[:4000]
    scores = {"es": 0, "en": 0, "pt": 0, "gl": 0}

    stopwords = {
        "es": [" el ", " la ", " de ", " que ", " y ", " en ", " los ", " las ", " para ", " con ", " una ", " por ", " del "],
        "en": [" the ", " and ", " of ", " to ", " in ", " for ", " with ", " on ", " this ", " that ", " is ", " are "],
        "pt": [" o ", " a ", " de ", " que ", " e ", " em ", " para ", " com ", " uma ", " os ", " as ", " do ", " da ", " não "],
        "gl": [" o ", " a ", " de ", " que ", " e ", " en ", " para ", " con ", " unha ", " os ", " as ", " do ", " da ", " non "],
    }

    padded = f" {sample} "
    for lang, words in stopwords.items():
        for word in words:
            scores[lang] += padded.count(word)

    if any(ch in sample for ch in "ñáéíóú¿¡"):
        scores["es"] += 2
    if any(ch in sample for ch in "ãõêçâô"):
        scores["pt"] += 3
    if " não " in padded:
        scores["pt"] += 3
    if " non " in padded:
        scores["gl"] += 3
    if " unha " in padded or " dende " in padded or " co " in padded:
        scores["gl"] += 2
    if " the " in padded or " this " in padded or " with " in padded:
        scores["en"] += 2

    return max(scores, key=scores.get)



def safe_round(value: Optional[float], digits: int = 2) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(float(value), digits)
    except Exception:
        return None



def count_words(text: str) -> int:
    return len(re.findall(r"\b\w+\b", text, flags=re.UNICODE))



def split_sentences(text: str) -> List[str]:
    text = normalize_whitespace(text)
    if not text:
        return []
    sentences = re.split(r"(?<=[\.\!\?])\s+", text)
    return [normalize_whitespace(s) for s in sentences if normalize_whitespace(s)]



def make_excerpt(text: str, max_chars: int = 220) -> str:
    text = normalize_whitespace(text)
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip() + "…"



def clamp(value: float, lower: float, upper: float) -> float:
    return max(lower, min(upper, value))



def normalize_bbox(bbox: Any) -> Optional[List[float]]:
    if bbox is None:
        return None

    candidate = bbox
    if isinstance(candidate, dict) and "bbox" in candidate:
        candidate = candidate.get("bbox")

    if isinstance(candidate, (list, tuple)) and len(candidate) == 4:
        raw_values = list(candidate)
    elif all(hasattr(candidate, attr) for attr in ("x0", "y0", "x1", "y1")):
        raw_values = [getattr(candidate, "x0"), getattr(candidate, "y0"), getattr(candidate, "x1"), getattr(candidate, "y1")]
    elif all(hasattr(candidate, attr) for attr in ("x", "y", "width", "height")):
        x = getattr(candidate, "x")
        y = getattr(candidate, "y")
        width = getattr(candidate, "width")
        height = getattr(candidate, "height")
        raw_values = [x, y, x + width, y + height]
    else:
        return None

    try:
        x0, y0, x1, y1 = [float(v) for v in raw_values]
        return [round(x0, 2), round(y0, 2), round(x1, 2), round(y1, 2)]
    except Exception:
        return None



def format_bbox(bbox: Optional[List[float]]) -> str:
    if not bbox:
        return "sin coordenadas"
    x0, y0, x1, y1 = bbox
    return f"x={x0:.1f}, y={y0:.1f}, w={(x1 - x0):.1f}, h={(y1 - y0):.1f}"


def humanize_bbox(bbox: Optional[List[float]], page_rect: Optional[List[float]] = None) -> str:
    if not bbox:
        return "ubicación no disponible"

    if page_rect and len(page_rect) == 4:
        try:
            px0, py0, px1, py1 = page_rect
            width = max(float(px1) - float(px0), 1.0)
            height = max(float(py1) - float(py0), 1.0)
            x0, y0, x1, y1 = bbox
            cx = ((float(x0) + float(x1)) / 2.0 - float(px0)) / width
            cy = ((float(y0) + float(y1)) / 2.0 - float(py0)) / height

            if cx < 0.33:
                hz = "izquierda"
            elif cx > 0.67:
                hz = "derecha"
            else:
                hz = "centro"

            if cy < 0.33:
                vt = "superior"
            elif cy > 0.67:
                vt = "inferior"
            else:
                vt = "central"

            mapping = {
                ("superior", "izquierda"): "zona superior izquierda",
                ("superior", "centro"): "zona superior central",
                ("superior", "derecha"): "zona superior derecha",
                ("central", "izquierda"): "zona media izquierda",
                ("central", "centro"): "zona central",
                ("central", "derecha"): "zona media derecha",
                ("inferior", "izquierda"): "zona inferior izquierda",
                ("inferior", "centro"): "zona inferior central",
                ("inferior", "derecha"): "zona inferior derecha",
            }
            return mapping.get((vt, hz), "zona aproximada del documento")
        except Exception:
            pass

    return "ubicación aproximada disponible"


def clean_list_items(items: Iterable[Any]) -> List[str]:
    out: List[str] = []
    for item in items:
        value = normalize_whitespace(str(item or ""))
        if value:
            out.append(value)
    return out


def issue_sort_key(issue: dict) -> Tuple[int, int, str, int]:
    severity_rank = {"high": 0, "medium": 1, "low": 2}.get(issue.get("severity_code"), 3)
    category_rank = {
        "text_contrast": 0,
        "figure_contrast": 1,
        "cvd_risk": 2,
        "dense_text": 3,
        "small_fonts": 4,
        "no_text": 5,
        "processing_note": 9,
    }.get(issue.get("category_code"), 8)
    unit_index = int(issue.get("unit_index", 0) or 0)
    return (severity_rank, category_rank, str(issue.get("location", "")), unit_index)


def classify_text_metric_status(ratio: Optional[float], threshold: Optional[float]) -> str:
    if ratio is None or threshold is None:
        return "sin dato"
    if ratio < threshold:
        return "incidencia"
    if ratio < threshold + 1.0:
        return "cercano al umbral"
    return "correcto"


def classify_figure_metric_status(ratio: Optional[float], threshold: Optional[float], worst_cvd_ratio: Optional[float]) -> str:
    if ratio is not None and threshold is not None and ratio < threshold:
        return "incidencia"
    if worst_cvd_ratio is not None and worst_cvd_ratio < 0.75:
        return "riesgo CVD"
    if ratio is not None and threshold is not None and ratio < threshold + 0.8:
        return "cercano al umbral"
    return "correcto"


def summarize_text_metrics(rows: List[dict]) -> dict:
    ratios = [float(row["ratio"]) for row in rows if row.get("ratio") is not None]
    estimated_count = sum(1 for row in rows if row.get("estimated"))
    below_count = sum(
        1
        for row in rows
        if row.get("ratio") is not None and row.get("threshold") is not None and float(row["ratio"]) < float(row["threshold"])
    )
    near_count = sum(
        1
        for row in rows
        if row.get("ratio") is not None
        and row.get("threshold") is not None
        and float(row["threshold"]) <= float(row["ratio"]) < float(row["threshold"]) + 1.0
    )
    return {
        "min_ratio": safe_round(min(ratios), 2) if ratios else None,
        "avg_ratio": safe_round(sum(ratios) / len(ratios), 2) if ratios else None,
        "estimated_count": estimated_count,
        "below_threshold_count": below_count,
        "near_threshold_count": near_count,
    }


def summarize_figure_metrics(rows: List[dict]) -> dict:
    ratios = [float(row["ratio"]) for row in rows if row.get("ratio") is not None]
    worst_scores = [float(row["worst_cvd_ratio"]) for row in rows if row.get("worst_cvd_ratio") is not None]
    low_contrast_count = sum(
        1
        for row in rows
        if row.get("ratio") is not None and row.get("threshold") is not None and float(row["ratio"]) < float(row["threshold"])
    )
    cvd_risk_count = sum(1 for row in rows if row.get("worst_cvd_ratio") is not None and float(row["worst_cvd_ratio"]) < 0.75)
    return {
        "min_ratio": safe_round(min(ratios), 2) if ratios else None,
        "avg_ratio": safe_round(sum(ratios) / len(ratios), 2) if ratios else None,
        "min_cvd_ratio": safe_round(min(worst_scores), 3) if worst_scores else None,
        "low_contrast_count": low_contrast_count,
        "cvd_risk_count": cvd_risk_count,
    }


def select_text_metric_rows(rows: List[dict], max_rows: int = 8) -> List[dict]:
    if not rows:
        return []
    scored = []
    for row in rows:
        ratio = float(row.get("ratio") or 0.0)
        threshold = float(row.get("threshold") or 0.0)
        delta = ratio - threshold
        status = classify_text_metric_status(row.get("ratio"), row.get("threshold"))
        scored.append((0 if ratio < threshold else 1 if delta < 1.0 else 2, delta, ratio, row.get("location", ""), row, status))
    scored.sort(key=lambda item: (item[0], item[1], item[2], item[3]))

    selected: List[dict] = []
    seen = set()
    for _bucket, _delta, _ratio, _location, row, status in scored:
        key = (row.get("location"), row.get("text_block_index"))
        if key in seen:
            continue
        seen.add(key)
        new_row = dict(row)
        new_row["status"] = status
        selected.append(new_row)
        if len(selected) >= max_rows:
            break
    return selected


def select_figure_metric_rows(rows: List[dict], max_rows: int = 8) -> List[dict]:
    if not rows:
        return []
    scored = []
    for row in rows:
        ratio = float(row.get("ratio") or 99.0) if row.get("ratio") is not None else 99.0
        threshold = float(row.get("threshold") or WCAG_NON_TEXT_RATIO)
        worst_cvd = float(row.get("worst_cvd_ratio") or 1.0) if row.get("worst_cvd_ratio") is not None else 1.0
        status = classify_figure_metric_status(row.get("ratio"), row.get("threshold"), row.get("worst_cvd_ratio"))
        bucket = 0 if (row.get("ratio") is not None and float(row.get("ratio")) < threshold) or worst_cvd < 0.75 else 1 if ratio < threshold + 0.8 or worst_cvd < 0.9 else 2
        scored.append((bucket, worst_cvd, ratio, row.get("location", ""), row, status))
    scored.sort(key=lambda item: (item[0], item[1], item[2], item[3]))

    selected: List[dict] = []
    seen = set()
    for _bucket, _worst_cvd, _ratio, _location, row, status in scored:
        key = (row.get("location"), row.get("figure_index"))
        if key in seen:
            continue
        seen.add(key)
        new_row = dict(row)
        new_row["status"] = status
        selected.append(new_row)
        if len(selected) >= max_rows:
            break
    return selected


def build_metric_summary_lines(summary: dict, kind: str) -> List[str]:
    if not summary:
        return []
    if kind == "text":
        return [
            f"Mínimo observado: {summary['min_ratio']}:1" if summary.get("min_ratio") is not None else "Mínimo observado: -",
            f"Media observada: {summary['avg_ratio']}:1" if summary.get("avg_ratio") is not None else "Media observada: -",
            f"Por debajo del umbral: {summary.get('below_threshold_count', 0)}",
            f"Cercanos al umbral: {summary.get('near_threshold_count', 0)}",
            f"Estimados por rasterización: {summary.get('estimated_count', 0)}",
        ]
    return [
        f"Mínimo de contraste observado: {summary['min_ratio']}:1" if summary.get("min_ratio") is not None else "Mínimo de contraste observado: -",
        f"Media observada: {summary['avg_ratio']}:1" if summary.get("avg_ratio") is not None else "Media observada: -",
        f"Figuras bajo umbral: {summary.get('low_contrast_count', 0)}",
        f"Peor score CVD: {summary['min_cvd_ratio']}" if summary.get("min_cvd_ratio") is not None else "Peor score CVD: -",
        f"Figuras con riesgo CVD: {summary.get('cvd_risk_count', 0)}",
    ]

def rgb_to_hex(rgb: Optional[Tuple[int, int, int]]) -> Optional[str]:
    if not rgb:
        return None
    try:
        r, g, b = [int(clamp(float(c), 0, 255)) for c in rgb]
        return f"#{r:02X}{g:02X}{b:02X}"
    except Exception:
        return None



def normalize_rgb(value: Any) -> Optional[Tuple[int, int, int]]:
    if value is None:
        return None
    if isinstance(value, tuple) and len(value) == 3:
        try:
            return tuple(int(clamp(float(c), 0, 255)) for c in value)
        except Exception:
            return None
    if hasattr(value, "rgb") and getattr(value, "rgb", None) is not None:
        value = getattr(value, "rgb")
    if hasattr(value, "__iter__") and not isinstance(value, (str, bytes)):
        data = list(value)
        if len(data) >= 3:
            try:
                return tuple(int(clamp(float(c), 0, 255)) for c in data[:3])
            except Exception:
                return None
    if isinstance(value, int):
        try:
            return ((value >> 16) & 255, (value >> 8) & 255, value & 255)
        except Exception:
            return None
    if isinstance(value, str):
        val = value.strip().lstrip("#")
        if len(val) == 6:
            try:
                return (int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16))
            except Exception:
                return None
    return None



def int_to_rgb(color_int: Any) -> Optional[Tuple[int, int, int]]:
    try:
        color_int = int(color_int)
    except Exception:
        return None
    return ((color_int >> 16) & 255, (color_int >> 8) & 255, color_int & 255)



def srgb_to_linear_channel(channel: float) -> float:
    channel = clamp(channel / 255.0, 0.0, 1.0)
    if channel <= 0.04045:
        return channel / 12.92
    return ((channel + 0.055) / 1.055) ** 2.4



def relative_luminance(rgb: Tuple[int, int, int]) -> float:
    r, g, b = rgb
    r_lin = srgb_to_linear_channel(r)
    g_lin = srgb_to_linear_channel(g)
    b_lin = srgb_to_linear_channel(b)
    return 0.2126 * r_lin + 0.7152 * g_lin + 0.0722 * b_lin



def contrast_ratio(rgb_a: Tuple[int, int, int], rgb_b: Tuple[int, int, int]) -> float:
    la = relative_luminance(rgb_a)
    lb = relative_luminance(rgb_b)
    lighter = max(la, lb)
    darker = min(la, lb)
    return (lighter + 0.05) / (darker + 0.05)



def infer_severity_for_ratio(ratio: float, threshold: float) -> str:
    if ratio >= threshold:
        return "low"
    if ratio >= threshold * 0.8:
        return "medium"
    return "high"



def severity_label(code: str) -> str:
    return {"low": "baja", "medium": "media", "high": "alta"}.get(code, code)



def is_large_text(size_pt: Optional[float]) -> bool:
    return bool(size_pt is not None and size_pt >= SMALL_FONT_THRESHOLD_PT)



def average_color(colors: List[Tuple[int, int, int]], weights: Optional[List[float]] = None) -> Optional[Tuple[int, int, int]]:
    if not colors:
        return None
    if weights and len(weights) == len(colors):
        total = float(sum(max(w, 0.0) for w in weights)) or 1.0
        r = sum(c[0] * max(w, 0.0) for c, w in zip(colors, weights)) / total
        g = sum(c[1] * max(w, 0.0) for c, w in zip(colors, weights)) / total
        b = sum(c[2] * max(w, 0.0) for c, w in zip(colors, weights)) / total
    else:
        r = sum(c[0] for c in colors) / len(colors)
        g = sum(c[1] for c in colors) / len(colors)
        b = sum(c[2] for c in colors) / len(colors)
    return (int(round(r)), int(round(g)), int(round(b)))



def get_optional_numpy():
    try:
        import numpy as np  # type: ignore

        return np
    except Exception:
        return None



def get_optional_pil_image():
    try:
        from PIL import Image  # type: ignore

        return Image
    except Exception:
        return None



def scale_image_if_needed(image: Any, max_side: int = MAX_RENDER_SIDE) -> Any:
    width, height = getattr(image, "size", (0, 0))
    if not width or not height:
        return image
    current_max = max(width, height)
    if current_max <= max_side:
        return image
    ratio = max_side / float(current_max)
    target = (max(1, int(width * ratio)), max(1, int(height * ratio)))
    try:
        return image.resize(target)
    except Exception:
        return image



def pil_image_from_bytes(blob: Optional[bytes]) -> Any:
    if not blob:
        return None
    Image = get_optional_pil_image()
    if Image is None:
        return None
    try:
        image = Image.open(io.BytesIO(blob))
        image.load()
        if image.mode != "RGB":
            image = image.convert("RGB")
        return scale_image_if_needed(image)
    except Exception:
        return None



def pil_image_from_pixmap(pixmap: Any) -> Any:
    Image = get_optional_pil_image()
    if Image is None or pixmap is None:
        return None
    try:
        mode = "RGBA" if getattr(pixmap, "alpha", False) else "RGB"
        image = Image.frombytes(mode, [pixmap.width, pixmap.height], pixmap.samples)
        if image.mode != "RGB":
            image = image.convert("RGB")
        return scale_image_if_needed(image)
    except Exception:
        return None



def image_to_palette(image: Any, color_count: int = 6) -> List[dict]:
    Image = get_optional_pil_image()
    np = get_optional_numpy()
    if image is None or Image is None or np is None:
        return []

    try:
        small = image.convert("RGB")
        small.thumbnail((180, 180))
        quantized = small.quantize(colors=color_count, method=Image.MEDIANCUT)
        palette = quantized.getpalette() or []
        color_counts = quantized.getcolors() or []
        total = float(sum(count for count, _ in color_counts) or 1)
        out: List[dict] = []
        for count, color_idx in sorted(color_counts, reverse=True):
            pos = color_idx * 3
            rgb = tuple(int(v) for v in palette[pos:pos + 3])
            if len(rgb) == 3:
                out.append(
                    {
                        "rgb": rgb,
                        "hex": rgb_to_hex(rgb),
                        "weight": round(count / total, 4),
                    }
                )
        return out[:color_count]
    except Exception:
        return []



def simulate_cvd_rgb(rgb: Tuple[int, int, int], deficiency: str) -> Tuple[int, int, int]:
    matrix = CVD_MATRICES.get(deficiency)
    if matrix is None:
        return rgb
    r, g, b = [c / 255.0 for c in rgb]
    r2 = matrix[0][0] * r + matrix[0][1] * g + matrix[0][2] * b
    g2 = matrix[1][0] * r + matrix[1][1] * g + matrix[1][2] * b
    b2 = matrix[2][0] * r + matrix[2][1] * g + matrix[2][2] * b
    return (
        int(round(clamp(r2, 0.0, 1.0) * 255)),
        int(round(clamp(g2, 0.0, 1.0) * 255)),
        int(round(clamp(b2, 0.0, 1.0) * 255)),
    )



def weighted_average_pairwise_distance(palette: List[dict], deficiency: Optional[str] = None) -> float:
    if len(palette) < 2:
        return 0.0
    total_weight = 0.0
    total_distance = 0.0
    for i in range(len(palette)):
        rgb_i = palette[i]["rgb"]
        if deficiency:
            rgb_i = simulate_cvd_rgb(rgb_i, deficiency)
        for j in range(i + 1, len(palette)):
            rgb_j = palette[j]["rgb"]
            if deficiency:
                rgb_j = simulate_cvd_rgb(rgb_j, deficiency)
            weight = max(float(palette[i].get("weight", 0.0)), 0.0) * max(float(palette[j].get("weight", 0.0)), 0.0)
            if weight <= 0:
                continue
            distance = math.sqrt(
                (rgb_i[0] - rgb_j[0]) ** 2 + (rgb_i[1] - rgb_j[1]) ** 2 + (rgb_i[2] - rgb_j[2]) ** 2
            )
            total_distance += distance * weight
            total_weight += weight
    if total_weight <= 0:
        return 0.0
    return total_distance / total_weight



def estimate_clip_contrast(image: Any) -> Optional[dict]:
    np = get_optional_numpy()
    if image is None or np is None:
        return None
    try:
        image = scale_image_if_needed(image, max_side=240).convert("RGB")
        arr = np.asarray(image, dtype=np.uint8)
        flat = arr.reshape(-1, 3)
        if len(flat) < 16:
            return None
        luminance = np.array([relative_luminance((int(r), int(g), int(b))) for r, g, b in flat])
        if luminance.size == 0:
            return None
        lower_thr = float(np.quantile(luminance, 0.15))
        upper_thr = float(np.quantile(luminance, 0.85))
        darker_pixels = flat[luminance <= lower_thr]
        lighter_pixels = flat[luminance >= upper_thr]
        if len(darker_pixels) == 0 or len(lighter_pixels) == 0:
            return None
        fg = tuple(int(round(v)) for v in darker_pixels.mean(axis=0))
        bg = tuple(int(round(v)) for v in lighter_pixels.mean(axis=0))
        ratio = contrast_ratio(fg, bg)
        return {
            "foreground_rgb": fg,
            "background_rgb": bg,
            "contrast_ratio": ratio,
            "estimated": True,
        }
    except Exception:
        return None



def assess_text_block_contrast(text_block: dict, unit: dict) -> Optional[dict]:
    block_font = normalize_rgb(text_block.get("font_color"))
    block_bg = normalize_rgb(text_block.get("background_color"))
    size_pt = safe_round(text_block.get("size_pt"), 1)
    threshold = WCAG_LARGE_TEXT_RATIO if is_large_text(size_pt) else WCAG_NORMAL_TEXT_RATIO

    if block_font and block_bg:
        ratio = contrast_ratio(block_font, block_bg)
        return {
            "ratio": ratio,
            "threshold": threshold,
            "estimated": False,
            "foreground_rgb": block_font,
            "background_rgb": block_bg,
            "size_pt": size_pt,
        }

    render_meta = text_block.get("rendered_contrast") or {}
    fg = normalize_rgb(render_meta.get("foreground_rgb"))
    bg = normalize_rgb(render_meta.get("background_rgb"))
    ratio = render_meta.get("contrast_ratio")
    if ratio is not None and fg and bg:
        return {
            "ratio": float(ratio),
            "threshold": threshold,
            "estimated": True,
            "foreground_rgb": fg,
            "background_rgb": bg,
            "size_pt": size_pt,
        }
    return None



def analyze_figure_visuals(figure: dict) -> Optional[dict]:
    image = figure.get("image")
    if image is None:
        return None
    palette = image_to_palette(image, color_count=6)
    clip_contrast = estimate_clip_contrast(image)
    contrast_val = clip_contrast.get("contrast_ratio") if clip_contrast else None
    base_distance = weighted_average_pairwise_distance(palette)
    cvd_results: Dict[str, dict] = {}
    worst_ratio = 1.0
    worst_type = None
    for deficiency in ("protanopia", "deuteranopia", "tritanopia"):
        transformed_distance = weighted_average_pairwise_distance(palette, deficiency=deficiency)
        ratio = transformed_distance / base_distance if base_distance > 0 else 1.0
        worst_ratio = min(worst_ratio, ratio)
        if worst_type is None or ratio == worst_ratio:
            worst_type = deficiency
        if ratio < 0.55:
            risk_level = "high"
        elif ratio < 0.75:
            risk_level = "medium"
        else:
            risk_level = "low"
        cvd_results[deficiency] = {
            "score_ratio": safe_round(ratio, 3),
            "risk_code": risk_level,
            "risk": severity_label(risk_level),
        }

    non_text_severity = "low"
    if contrast_val is not None:
        non_text_severity = infer_severity_for_ratio(float(contrast_val), WCAG_NON_TEXT_RATIO)

    color_dependency = "moderada"
    if base_distance > 100 and worst_ratio < 0.7:
        color_dependency = "alta"
    elif base_distance < 45:
        color_dependency = "baja"

    return {
        "contrast_ratio": safe_round(contrast_val, 2),
        "non_text_threshold": WCAG_NON_TEXT_RATIO,
        "non_text_severity_code": non_text_severity,
        "non_text_severity": severity_label(non_text_severity),
        "palette": palette,
        "cvd": cvd_results,
        "worst_cvd_ratio": safe_round(worst_ratio, 3),
        "worst_cvd_type": worst_type,
        "color_dependency": color_dependency,
        "estimated": True,
    }



def render_pdf_block_contrast_estimates(unit: dict, logger: logging.Logger) -> None:
    page_ref = unit.get("_page_ref")
    fitz_module = unit.get("_fitz_module")
    if page_ref is None or fitz_module is None:
        return
    blocks = unit.get("text_blocks", []) or []
    if not blocks:
        return

    for block in blocks:
        bbox = block.get("bbox")
        if not bbox:
            continue
        try:
            rect = fitz_module.Rect(*bbox)
            matrix = fitz_module.Matrix(1.8, 1.8)
            pix = page_ref.get_pixmap(matrix=matrix, clip=rect, alpha=False)
            image = pil_image_from_pixmap(pix)
            if image is None:
                continue
            contrast_data = estimate_clip_contrast(image)
            if contrast_data:
                block["rendered_contrast"] = contrast_data
        except Exception as exc:
            logger.debug("No se pudo estimar el contraste renderizado del bloque %s: %s", bbox, exc)



def paragraph_lengths(text: str) -> List[int]:
    paragraphs = [p.strip() for p in re.split(r"\n+", normalize_whitespace(text)) if p.strip()]
    return [count_words(p) for p in paragraphs]



def friendly_units_label(document_type: str) -> str:
    return "páginas" if document_type == "pdf" else "diapositivas"


class NllbTranslator:
    def __init__(self, model_name: str = DEFAULT_TRANSLATION_MODEL) -> None:
        from transformers import AutoModelForSeq2SeqLM, AutoTokenizer
        import torch

        self.torch = torch

        # Detectar dispositivo óptimo
        if torch.cuda.is_available():
            self.device = "cuda"
            torch_threads = 0
        else:
            self.device = "cpu"
            torch_threads = _CPU_CORES

        self.tokenizer = AutoTokenizer.from_pretrained(model_name)
        self.model = AutoModelForSeq2SeqLM.from_pretrained(model_name)
        self.model.to(self.device)
        self.model.eval()

        if self.device == "cpu":
            try:
                torch.set_num_threads(torch_threads)
            except Exception:
                pass

    def translate_text(self, text: str, src_code: str, tgt_code: str, max_chars: int = 800) -> str:
        text = normalize_whitespace(text)
        if not text:
            return text
        if src_code == tgt_code:
            return text

        src_nllb = NLLB_LANG_MAP.get(src_code)
        tgt_nllb = NLLB_LANG_MAP.get(tgt_code)
        if not src_nllb or not tgt_nllb:
            raise ValueError(f"Idioma no soportado para traducción NLLB: {src_code} -> {tgt_code}")

        chunks = split_text_for_translation(text, max_chars=max_chars)
        out_chunks: List[str] = []

        for chunk in chunks:
            self.tokenizer.src_lang = src_nllb
            inputs = self.tokenizer(chunk, return_tensors="pt", truncation=True, padding=True)
            inputs = {k: v.to(self.device) for k, v in inputs.items()}
            with self.torch.no_grad():
                generated = self.model.generate(
                    **inputs,
                    forced_bos_token_id=self.tokenizer.convert_tokens_to_ids(tgt_nllb),
                    max_length=512,
                )
            translated = self.tokenizer.batch_decode(generated, skip_special_tokens=True)[0]
            out_chunks.append(normalize_whitespace(translated))

        return normalize_whitespace("\n\n".join(out_chunks))


# ---------------------------------------------------------
# EXTRACTION
# ---------------------------------------------------------


def extract_pdf_units(input_file: Path, use_ocr: bool, logger: logging.Logger) -> List[dict]:
    from pypdf import PdfReader

    logger.info("Extrayendo texto de PDF: %s", input_file.name)
    reader = PdfReader(str(input_file))
    units: List[dict] = []

    fitz_module = None
    try:
        import fitz as _fitz  # type: ignore

        fitz_module = _fitz
    except Exception:
        fitz_module = None

    fitz_doc = None
    if fitz_module:
        try:
            fitz_doc = fitz_module.open(str(input_file))
        except Exception:
            fitz_doc = None

    for idx, page in enumerate(reader.pages, start=1):
        text = normalize_whitespace(page.extract_text() or "")
        notes: List[str] = []
        text_blocks: List[dict] = []
        figures: List[dict] = []
        page_rect: Optional[List[float]] = None

        fitz_page = None
        if fitz_doc is not None:
            try:
                fitz_page = fitz_doc[idx - 1]
                page_rect = normalize_bbox(fitz_page.rect)
                page_dict = fitz_page.get_text("dict")
                structured_texts: List[str] = []
                for block_idx, block in enumerate(page_dict.get("blocks", []), start=1):
                    block_type = int(block.get("type", 0) or 0)
                    bbox = normalize_bbox(block.get("bbox"))
                    if block_type == 0:
                        block_texts: List[str] = []
                        span_sizes: List[float] = []
                        span_colors: List[Tuple[int, int, int]] = []
                        for line in block.get("lines", []):
                            for span in line.get("spans", []):
                                span_text = normalize_whitespace(span.get("text", ""))
                                if not span_text:
                                    continue
                                structured_texts.append(span_text)
                                block_texts.append(span_text)
                                size_val = safe_round(span.get("size"), 1)
                                if size_val is not None:
                                    span_sizes.append(size_val)
                                rgb = int_to_rgb(span.get("color"))
                                if rgb:
                                    span_colors.append(rgb)
                        joined = normalize_whitespace(" ".join(block_texts))
                        if joined:
                            text_blocks.append(
                                {
                                    "text_block_index": block_idx,
                                    "text": joined,
                                    "bbox": bbox,
                                    "size_pt": safe_round(sum(span_sizes) / len(span_sizes), 1) if span_sizes else None,
                                    "font_color": average_color(span_colors),
                                    "background_color": None,
                                    "source": "pdf_structured_block",
                                }
                            )
                    elif block_type == 1:
                        figures.append(
                            {
                                "figure_index": len(figures) + 1,
                                "location": f"Página {idx} · figura {len(figures) + 1}",
                                "bbox": bbox,
                                "image": None,
                                "notes": ["Figura detectada en PDF mediante bloque gráfico."],
                            }
                        )
                if not text and structured_texts:
                    text = normalize_whitespace("\n\n".join(structured_texts))
            except Exception as exc:
                logger.warning("No se pudo extraer estructura avanzada de la página %s: %s", idx, exc)

        if not text and fitz_doc is not None:
            try:
                text = normalize_whitespace(fitz_doc[idx - 1].get_text("text") or "")
            except Exception:
                pass

        if fitz_doc is not None and fitz_page is not None:
            try:
                seen_rects = set()
                for image_info in fitz_page.get_images(full=True):
                    xref = image_info[0]
                    rects = fitz_page.get_image_rects(xref) or []
                    image_blob = None
                    try:
                        extracted = fitz_doc.extract_image(xref)
                        image_blob = extracted.get("image") if extracted else None
                    except Exception:
                        image_blob = None
                    for rect in rects:
                        bbox = normalize_bbox(rect)
                        key = tuple(bbox or [])
                        if key in seen_rects:
                            continue
                        seen_rects.add(key)
                        figures.append(
                            {
                                "figure_index": len(figures) + 1,
                                "location": f"Página {idx} · figura {len(figures) + 1}",
                                "bbox": bbox,
                                "image": pil_image_from_bytes(image_blob),
                                "notes": [],
                            }
                        )
            except Exception as exc:
                notes.append(f"No se pudieron aislar todas las figuras incrustadas del PDF: {exc}")

        if fitz_doc is not None and fitz_page is not None and text_blocks:
            unit_stub = {"_page_ref": fitz_page, "_fitz_module": fitz_module, "text_blocks": text_blocks}
            render_pdf_block_contrast_estimates(unit_stub, logger=logger)

        if use_ocr and not text:
            notes.append("No se ha aplicado OCR porque no hay un motor OCR configurado en este pipeline.")

        units.append(
            {
                "unit_type": "page",
                "unit_index": idx,
                "label": f"Página {idx}",
                "title": f"Página {idx}",
                "text": text,
                "text_blocks": text_blocks,
                "figures": figures,
                "small_font_hits": [],
                "notes": notes,
                "page_rect": page_rect,
            }
        )

    return units



def _convert_ppt_to_pptx(input_file: Path, logger: logging.Logger) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "No se encontró LibreOffice/soffice para convertir .ppt a .pptx. "
            "Convierte el archivo a .pptx o instala LibreOffice en el servidor."
        )

    tmp_dir = Path(tempfile.mkdtemp(prefix="aiuda_ppt_convert_"))
    cmd = [
        soffice,
        "--headless",
        "--convert-to",
        "pptx",
        "--outdir",
        str(tmp_dir),
        str(input_file),
    ]
    logger.info("Convirtiendo PPT a PPTX con LibreOffice")
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError(f"No se pudo convertir el PPT: {proc.stderr.strip() or proc.stdout.strip()}")

    candidates = list(tmp_dir.glob("*.pptx"))
    if not candidates:
        raise RuntimeError("La conversión de PPT a PPTX no generó ningún archivo de salida.")
    return candidates[0]



def get_ppt_fill_rgb(fill: Any) -> Optional[Tuple[int, int, int]]:
    if fill is None:
        return None
    try:
        fore = getattr(fill, "fore_color", None)
        if fore is not None and getattr(fore, "rgb", None) is not None:
            return normalize_rgb(fore.rgb)
    except Exception:
        return None
    return None



def get_shape_background_color(shape: Any, default_bg: Optional[Tuple[int, int, int]]) -> Optional[Tuple[int, int, int]]:
    try:
        fill = getattr(shape, "fill", None)
        rgb = get_ppt_fill_rgb(fill)
        if rgb:
            return rgb
    except Exception:
        pass
    return default_bg



def get_run_font_color(run: Any, fallback: Optional[Tuple[int, int, int]]) -> Optional[Tuple[int, int, int]]:
    try:
        color = getattr(getattr(run, "font", None), "color", None)
        if color is not None and getattr(color, "rgb", None) is not None:
            rgb = normalize_rgb(color.rgb)
            if rgb:
                return rgb
    except Exception:
        pass
    return fallback



def extract_pptx_units(input_file: Path, logger: logging.Logger) -> List[dict]:
    from pptx import Presentation

    logger.info("Extrayendo texto de presentación: %s", input_file.name)
    prs = Presentation(str(input_file))
    units: List[dict] = []
    slide_width = float(getattr(prs, "slide_width", 0) or 0)
    slide_height = float(getattr(prs, "slide_height", 0) or 0)

    for idx, slide in enumerate(prs.slides, start=1):
        texts: List[str] = []
        title_text = ""
        small_font_hits: List[dict] = []
        text_blocks: List[dict] = []
        figures: List[dict] = []
        notes: List[str] = []

        slide_bg = (255, 255, 255)
        try:
            bg_fill = getattr(slide.background, "fill", None)
            slide_bg = get_ppt_fill_rgb(bg_fill) or slide_bg
        except Exception:
            slide_bg = (255, 255, 255)

        for shape_idx, shape in enumerate(slide.shapes, start=1):
            bbox = normalize_bbox([
                float(getattr(shape, "left", 0) or 0),
                float(getattr(shape, "top", 0) or 0),
                float((getattr(shape, "left", 0) or 0) + (getattr(shape, "width", 0) or 0)),
                float((getattr(shape, "top", 0) or 0) + (getattr(shape, "height", 0) or 0)),
            ])

            if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                raw_text = normalize_whitespace(shape.text_frame.text)
                if raw_text:
                    texts.append(raw_text)
                try:
                    if slide.shapes.title is shape and raw_text:
                        title_text = raw_text
                except Exception:
                    pass

                block_bg = get_shape_background_color(shape, slide_bg)
                run_sizes: List[float] = []
                run_colors: List[Tuple[int, int, int]] = []
                block_text_parts: List[str] = []

                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run_text = normalize_whitespace(run.text)
                        if not run_text:
                            continue
                        block_text_parts.append(run_text)
                        size = getattr(run.font, "size", None)
                        size_pt = None
                        try:
                            if size is not None:
                                size_pt = float(size.pt)
                        except Exception:
                            size_pt = None
                        if size_pt is not None:
                            run_sizes.append(size_pt)
                        run_color = get_run_font_color(run, fallback=None)
                        if run_color:
                            run_colors.append(run_color)
                        if run_text and size_pt is not None and size_pt < SMALL_FONT_THRESHOLD_PT:
                            small_font_hits.append(
                                {
                                    "text": run_text[:120],
                                    "size_pt": round(size_pt, 1),
                                }
                            )

                block_text = normalize_whitespace(" ".join(block_text_parts)) or raw_text
                if block_text:
                    text_blocks.append(
                        {
                            "text_block_index": len(text_blocks) + 1,
                            "text": block_text,
                            "bbox": bbox,
                            "size_pt": safe_round(sum(run_sizes) / len(run_sizes), 1) if run_sizes else None,
                            "font_color": average_color(run_colors),
                            "background_color": block_bg,
                            "source": "pptx_text_frame",
                        }
                    )

            try:
                if hasattr(shape, "image") and getattr(shape, "image", None) is not None:
                    figures.append(
                        {
                            "figure_index": len(figures) + 1,
                            "location": f"Diapositiva {idx} · figura {len(figures) + 1}",
                            "bbox": bbox,
                            "image": pil_image_from_bytes(shape.image.blob),
                            "notes": [],
                        }
                    )
            except Exception as exc:
                notes.append(f"No se pudo procesar una figura de la diapositiva {idx}: {exc}")

        text = normalize_whitespace("\n\n".join(texts))
        units.append(
            {
                "unit_type": "slide",
                "unit_index": idx,
                "label": f"Diapositiva {idx}",
                "title": title_text or f"Diapositiva {idx}",
                "text": text,
                "text_blocks": text_blocks,
                "figures": figures,
                "small_font_hits": small_font_hits,
                "notes": notes,
                "page_rect": normalize_bbox([0, 0, slide_width, slide_height]),
            }
        )

    return units



def extract_document_units(input_file: Path, use_ocr: bool, logger: logging.Logger) -> Tuple[str, List[dict]]:
    suffix = input_file.suffix.lower()

    if suffix == ".pdf":
        return "pdf", extract_pdf_units(input_file, use_ocr, logger)

    if suffix == ".pptx":
        return "pptx", extract_pptx_units(input_file, logger)

    if suffix == ".ppt":
        converted = _convert_ppt_to_pptx(input_file, logger)
        return "ppt", extract_pptx_units(converted, logger)

    raise RuntimeError("Formato documental no soportado. Usa PDF, PPT o PPTX.")


# ---------------------------------------------------------
# ANALYSIS
# ---------------------------------------------------------


def add_issue(
    issues: List[dict],
    issue_id: int,
    location: str,
    category_code: str,
    category: str,
    severity_code: str,
    severity: str,
    message: str,
    recommendation: str,
    excerpt: str,
    unit_index: int,
    extra: Optional[dict] = None,
) -> int:
    payload = {
        "id": f"issue-{issue_id:03d}",
        "location": location,
        "category_code": category_code,
        "category": category,
        "severity_code": severity_code,
        "severity": severity,
        "message": message,
        "recommendation": recommendation,
        "excerpt": excerpt,
        "unit_index": unit_index,
    }
    if extra:
        payload.update(extra)
    issues.append(payload)
    return issue_id + 1



def build_methodology_section() -> dict:
    return {
        "title": "Metodología y referencias",
        "contrast_rules": [
            {
                "metric_key": "text_normal",
                "label": "Contraste textual para texto normal",
                "metric_type_code": "normative",
                "metric_type": "Umbral normativo de referencia",
                "threshold": WCAG_NORMAL_TEXT_RATIO,
                "reference_code": "WCAG-1.4.3",
                "details": "Referencia orientativa: 4.5:1.",
            },
            {
                "metric_key": "text_large",
                "label": "Contraste textual para texto grande",
                "metric_type_code": "normative",
                "metric_type": "Umbral normativo de referencia",
                "threshold": WCAG_LARGE_TEXT_RATIO,
                "reference_code": "WCAG-1.4.3",
                "details": "Referencia orientativa: 3:1.",
            },
            {
                "metric_key": "non_text",
                "label": "Contraste de figuras y componentes visuales relevantes",
                "metric_type_code": "normative",
                "metric_type": "Umbral normativo de referencia",
                "threshold": WCAG_NON_TEXT_RATIO,
                "reference_code": "WCAG-1.4.11",
                "details": "Referencia orientativa: 3:1.",
            },
            {
                "metric_key": "cvd_risk",
                "label": "Riesgo CVD / daltonismo",
                "metric_type_code": "heuristic",
                "metric_type": "Métrica heurística del sistema",
                "threshold": None,
                "reference_code": "BRETTEL-1997 + MACHADO-2009",
                "details": "Estimación automática basada en simulaciones aproximadas de protanopia, deuteranopia y tritanopia. No equivale a certificación oficial.",
            },
        ],
        "references": deepcopy(METHODOLOGY_REFERENCES),
        "limitations": [
            "En PPT/PPTX el contraste puede calcularse con más precisión cuando existen colores RGB explícitos en shapes y runs.",
            "En PDF, cuando no hay colores estructurados, el contraste textual puede estimarse a partir de renderización local del bloque y debe interpretarse como una estimación.",
            "La evaluación CVD es heurística y orientativa: ayuda a detectar riesgo relativo de confusión cromática, pero no sustituye una revisión experta.",
        ],
    }



def analyze_document(units: List[dict], document_type: str, input_filename: str) -> dict:
    issues: List[dict] = []
    recommendations: List[str] = []
    issue_id = 1
    dense_text_count = 0
    small_font_count = 0
    text_contrast_checks = 0
    text_contrast_issue_count = 0
    figures_total = 0
    figures_analyzed = 0
    figure_contrast_issue_count = 0
    cvd_issue_count = 0
    processing_notes: List[str] = []
    all_text_metric_rows: List[dict] = []
    all_figure_metric_rows: List[dict] = []

    for unit in units:
        text = normalize_whitespace(unit.get("text", ""))
        location = unit.get("label", "Documento")
        unit_index = int(unit.get("unit_index", 0) or 0)
        excerpt = make_excerpt(text)
        text_blocks = unit.get("text_blocks", []) or []
        figures = unit.get("figures", []) or []
        page_rect = normalize_bbox(unit.get("page_rect"))
        figures_total += len(figures)

        if not text and not figures:
            issue_id = add_issue(
                issues,
                issue_id,
                location,
                "no_text",
                "Sin texto detectable",
                "low",
                "baja",
                "No se ha detectado texto utilizable en esta sección del documento.",
                "Verifica si la página o diapositiva contiene texto seleccionable o si requiere una revisión manual.",
                excerpt,
                unit_index,
            )

        if text:
            word_count = count_words(text)
            paragraph_word_lengths = paragraph_lengths(text)
            longest_paragraph_words = max(paragraph_word_lengths, default=0)
            if word_count >= TEXT_DENSITY_WORD_THRESHOLD or longest_paragraph_words >= PARAGRAPH_DENSITY_WORD_THRESHOLD:
                dense_text_count += 1
                issue_id = add_issue(
                    issues,
                    issue_id,
                    location,
                    "dense_text",
                    "Carga visual de texto",
                    "medium",
                    "media",
                    "Se ha detectado un bloque textual denso que puede dificultar el escaneo visual y la lectura rápida.",
                    "Divide la información en bloques más breves, usa más aire visual y redistribuye el contenido entre más páginas o diapositivas cuando sea necesario.",
                    excerpt,
                    unit_index,
                    extra={
                        "metric_type": "Métrica heurística del sistema",
                        "word_count": word_count,
                        "longest_paragraph_words": longest_paragraph_words,
                    },
                )

        if document_type in {"ppt", "pptx"} and unit.get("small_font_hits"):
            sample = unit["small_font_hits"][0]
            small_font_count += len(unit["small_font_hits"])
            issue_id = add_issue(
                issues,
                issue_id,
                location,
                "small_fonts",
                "Tamaño de letra",
                "medium",
                "media",
                f"Se han detectado textos potencialmente pequeños para una visualización cómoda. Ejemplo aproximado: {sample.get('size_pt', '?')} pt.",
                "Aumenta el tamaño de las fuentes y verifica la legibilidad en pantalla completa, proyección o móvil.",
                make_excerpt(sample.get("text", "")),
                unit_index,
                extra={
                    "metric_type": "Métrica heurística del sistema",
                    "threshold": SMALL_FONT_THRESHOLD_PT,
                },
            )

        for block in text_blocks:
            contrast_data = assess_text_block_contrast(block, unit)
            if not contrast_data:
                continue
            text_contrast_checks += 1
            ratio = float(contrast_data["ratio"])
            threshold = float(contrast_data["threshold"])
            estimated = bool(contrast_data.get("estimated"))
            bbox = normalize_bbox(block.get("bbox"))
            bbox_human = humanize_bbox(bbox, page_rect)
            all_text_metric_rows.append(
                {
                    "location": location,
                    "text_block_index": block.get("text_block_index"),
                    "ratio": safe_round(ratio, 2),
                    "threshold": threshold,
                    "estimated": estimated,
                    "method": "estimado" if estimated else "estructurado",
                    "size_pt": contrast_data.get("size_pt"),
                    "bbox": bbox,
                    "bbox_human": bbox_human,
                    "color_fg": rgb_to_hex(contrast_data.get("foreground_rgb")),
                    "color_bg": rgb_to_hex(contrast_data.get("background_rgb")),
                }
            )

            if ratio < threshold:
                text_contrast_issue_count += 1
                severity_code = infer_severity_for_ratio(ratio, threshold)
                issue_id = add_issue(
                    issues,
                    issue_id,
                    f"{location} · bloque textual {block.get('text_block_index', '?')}",
                    "text_contrast",
                    "Contraste textual",
                    severity_code,
                    severity_label(severity_code),
                    (
                        f"Se estima un contraste texto-fondo de {ratio:.2f}:1, por debajo del umbral de referencia {threshold:.1f}:1. "
                        f"Zona aproximada: {bbox_human}."
                    ),
                    "Aumenta la diferencia entre color de texto y fondo, o reestructura el bloque visual para cumplir mejor con WCAG/WCAG2ICT.",
                    make_excerpt(block.get("text", "")),
                    unit_index,
                    extra={
                        "metric_type": "Umbral normativo de referencia" if not estimated else "Métrica heurística del sistema",
                        "ratio": safe_round(ratio, 2),
                        "threshold": threshold,
                        "bbox": bbox,
                        "bbox_human": bbox_human,
                        "color_fg": rgb_to_hex(contrast_data.get("foreground_rgb")),
                        "color_bg": rgb_to_hex(contrast_data.get("background_rgb")),
                        "estimated": estimated,
                        "reference_code": "WCAG-1.4.3 / WCAG2ICT",
                    },
                )

        for figure in figures:
            figure_metrics = analyze_figure_visuals(figure)
            if not figure_metrics:
                note = f"{figure.get('location', location)}: no se pudo calcular la métrica visual figura a figura con fiabilidad suficiente."
                processing_notes.append(note)
                continue

            figures_analyzed += 1
            ratio = figure_metrics.get("contrast_ratio")
            bbox = normalize_bbox(figure.get("bbox"))
            bbox_human = humanize_bbox(bbox, page_rect)
            all_figure_metric_rows.append(
                {
                    "location": figure.get("location", location),
                    "figure_index": figure.get("figure_index"),
                    "ratio": ratio,
                    "threshold": WCAG_NON_TEXT_RATIO,
                    "bbox": bbox,
                    "bbox_human": bbox_human,
                    "worst_cvd_type": figure_metrics.get("worst_cvd_type"),
                    "worst_cvd_ratio": figure_metrics.get("worst_cvd_ratio"),
                    "color_dependency": figure_metrics.get("color_dependency"),
                }
            )

            if ratio is not None and float(ratio) < WCAG_NON_TEXT_RATIO:
                figure_contrast_issue_count += 1
                severity_code = infer_severity_for_ratio(float(ratio), WCAG_NON_TEXT_RATIO)
                issue_id = add_issue(
                    issues,
                    issue_id,
                    figure.get("location", location),
                    "figure_contrast",
                    "Contraste en figura",
                    severity_code,
                    severity_label(severity_code),
                    (
                        f"La figura presenta un contraste visual estimado de {float(ratio):.2f}:1, por debajo del umbral orientativo {WCAG_NON_TEXT_RATIO:.1f}:1 para elementos no textuales relevantes. "
                        f"Zona aproximada: {bbox_human}."
                    ),
                    "Revisa si la figura depende de diferencias muy sutiles de luminancia o color, y refuerza contornos, etiquetas o contraste.",
                    f"Ubicación aproximada: {bbox_human}",
                    unit_index,
                    extra={
                        "metric_type": "Métrica heurística del sistema",
                        "ratio": ratio,
                        "threshold": WCAG_NON_TEXT_RATIO,
                        "bbox": bbox,
                        "bbox_human": bbox_human,
                        "reference_code": "WCAG-1.4.11 / WCAG2ICT",
                    },
                )

            worst_ratio = figure_metrics.get("worst_cvd_ratio")
            if worst_ratio is not None and float(worst_ratio) < 0.75:
                cvd_issue_count += 1
                severity_code = "high" if float(worst_ratio) < 0.55 else "medium"
                issue_id = add_issue(
                    issues,
                    issue_id,
                    figure.get("location", location),
                    "cvd_risk",
                    "Riesgo para daltonismo",
                    severity_code,
                    severity_label(severity_code),
                    (
                        f"La figura muestra una pérdida relativa de separabilidad cromática bajo simulación CVD. El caso más sensible es {figure_metrics.get('worst_cvd_type', 'desconocido')} con score {float(worst_ratio):.3f}."
                    ),
                    "Añade codificación redundante mediante etiquetas, patrones, contornos o diferencias de luminancia para reducir la dependencia exclusiva del color.",
                    f"Zona aproximada: {bbox_human}. Dependencia del color estimada: {figure_metrics.get('color_dependency', 'desconocida')}",
                    unit_index,
                    extra={
                        "metric_type": "Métrica heurística del sistema",
                        "score_ratio": worst_ratio,
                        "reference_code": "BRETTEL-1997 / MACHADO-2009",
                        "bbox": bbox,
                        "bbox_human": bbox_human,
                        "cvd": figure_metrics.get("cvd"),
                    },
                )

        for note in unit.get("notes", []):
            clean_note = normalize_whitespace(note)
            if not clean_note:
                continue
            processing_notes.append(f"{location}: {clean_note}")
            issue_id = add_issue(
                issues,
                issue_id,
                location,
                "processing_note",
                "Nota de procesamiento",
                "low",
                "baja",
                clean_note,
                "Revisa manualmente esta sección del documento para completar la validación.",
                excerpt,
                unit_index,
                extra={"metric_type": "Métrica heurística del sistema"},
            )

    if dense_text_count:
        recommendations.append(
            "Reduce la densidad de contenido por página o diapositiva, priorizando bloques más cortos, jerarquía visual más clara y más espacio en blanco."
        )
    if small_font_count:
        recommendations.append(
            "Revisa el tamaño de las fuentes y asegúrate de que sigan siendo legibles en pantalla compartida, proyección o visualización móvil."
        )
    if text_contrast_issue_count:
        recommendations.append(
            "Ajusta pares texto-fondo con bajo contraste, especialmente en bloques pequeños, subtítulos o etiquetas con colores suaves."
        )
    if figure_contrast_issue_count:
        recommendations.append(
            "Refuerza el contraste de figuras y gráficos con contornos, etiquetas directas y diferencias de luminancia más claras."
        )
    if cvd_issue_count:
        recommendations.append(
            "Evita depender solo del color en figuras relevantes: añade texto, patrones, marcadores o codificación redundante accesible para personas con CVD."
        )
    if not recommendations:
        recommendations.append(
            "El análisis automático no detectó alertas visuales relevantes, aunque siempre conviene una revisión manual final de contraste, jerarquía visual y dependencia del color."
        )

    recommendations = clean_list_items(recommendations)
    processing_notes = clean_list_items(list(dict.fromkeys(processing_notes)))[:20]
    issues.sort(key=issue_sort_key)
    reportable_issues = [issue for issue in issues if issue.get("category_code") != "processing_note"]

    units_label = friendly_units_label(document_type)
    overview = (
        f"Se analizaron {len(units)} {units_label} del archivo {input_filename}. "
        f"Se detectaron {len(reportable_issues)} alertas visuales y {len(processing_notes)} notas de procesamiento relevantes."
    )

    severity_totals = {
        "alta": sum(1 for issue in reportable_issues if issue.get("severity_code") == "high"),
        "media": sum(1 for issue in reportable_issues if issue.get("severity_code") == "medium"),
        "baja": sum(1 for issue in reportable_issues if issue.get("severity_code") == "low"),
    }

    methodology = build_methodology_section()
    text_metric_samples = select_text_metric_rows(all_text_metric_rows, max_rows=8)
    figure_metric_samples = select_figure_metric_rows(all_figure_metric_rows, max_rows=8)
    visual_metrics = {
        "reportable_issue_count": len(reportable_issues),
        "processing_note_count": len(processing_notes),
        "text_blocks_analyzed": text_contrast_checks,
        "text_contrast_issues": text_contrast_issue_count,
        "dense_text_sections": dense_text_count,
        "small_font_hits": small_font_count,
        "figure_count": figures_total,
        "figures_analyzed": figures_analyzed,
        "figure_contrast_issues": figure_contrast_issue_count,
        "cvd_risk_issues": cvd_issue_count,
        "text_metrics_summary": summarize_text_metrics(all_text_metric_rows),
        "figure_metrics_summary": summarize_figure_metrics(all_figure_metric_rows),
        "text_metric_samples": text_metric_samples,
        "figure_metric_samples": figure_metric_samples,
        "processing_notes": processing_notes,
    }

    return {
        "summary": {
            "title": "Informe de accesibilidad documental",
            "overview": overview,
            "totals": {
                "units_analyzed": len(units),
                "issue_count": len(issues),
                "reportable_issue_count": len(reportable_issues),
                "processing_note_count": len(processing_notes),
                "document_type": document_type,
                "text_blocks_analyzed": text_contrast_checks,
                "figures_analyzed": figures_analyzed,
            },
            "severity_totals": severity_totals,
        },
        "issues": issues,
        "recommendations": recommendations,
        "visual_metrics": visual_metrics,
        "methodology": methodology,
    }

def translate_payload_strings(
    obj: Any,
    translator: Optional[NllbTranslator],
    tgt_lang: str,
    logger: logging.Logger,
    max_chars: int,
    current_key: Optional[str] = None,
) -> Any:
    if isinstance(obj, dict):
        out = {}
        for key, value in obj.items():
            out[key] = translate_payload_strings(
                value,
                translator=translator,
                tgt_lang=tgt_lang,
                logger=logger,
                max_chars=max_chars,
                current_key=key,
            )
        return out

    if isinstance(obj, list):
        return [
            translate_payload_strings(
                item,
                translator=translator,
                tgt_lang=tgt_lang,
                logger=logger,
                max_chars=max_chars,
                current_key=current_key,
            )
            for item in obj
        ]

    if isinstance(obj, str):
        text = normalize_whitespace(obj)
        if not text or current_key in SKIP_TRANSLATION_KEYS:
            return obj
        if tgt_lang == "es":
            return obj
        if translator is None:
            return obj
        try:
            return translator.translate_text(text, src_code="es", tgt_code=tgt_lang, max_chars=max_chars)
        except Exception as exc:
            logger.warning("No se pudo traducir el campo '%s' a %s: %s", current_key, tgt_lang, exc)
            return obj

    return obj


# ---------------------------------------------------------
# OUTPUTS
# ---------------------------------------------------------


def render_metric_value(value: Any) -> str:
    if value is None:
        return "-"
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)



def _filter_reportable_issues(payload: dict) -> List[dict]:
    all_issues = payload.get("issues", []) or []
    issues = [issue for issue in all_issues if issue.get("category_code") != "processing_note"]
    return sorted(issues, key=issue_sort_key)



def _teacher_area_label(category_code: str) -> str:
    mapping = {
        "cvd_risk": "Color, daltonismo y figuras",
        "figure_contrast": "Color, daltonismo y figuras",
        "text_contrast": "Tamaño de letra y legibilidad",
        "small_fonts": "Tamaño de letra y legibilidad",
        "dense_text": "Cantidad de texto",
        "no_text": "Contenido no detectable",
    }
    return mapping.get(normalize_whitespace(str(category_code or "")).lower(), "Accesibilidad visual")



def _teacher_area_key(category_code: str) -> str:
    category_code = normalize_whitespace(str(category_code or "")).lower()
    if category_code in {"cvd_risk", "figure_contrast"}:
        return "color"
    if category_code in {"text_contrast", "small_fonts"}:
        return "legibility"
    if category_code == "dense_text":
        return "text_load"
    if category_code == "no_text":
        return "content"
    return "other"



def _teacher_severity_weight(issue: dict) -> float:
    severity_code = normalize_whitespace(str(issue.get("severity_code") or "")).lower()
    base = {"high": 3.0, "medium": 2.0, "low": 1.0}.get(severity_code, 1.0)
    category_code = normalize_whitespace(str(issue.get("category_code") or "")).lower()
    bonus = {
        "cvd_risk": 1.0,
        "figure_contrast": 0.7,
        "small_fonts": 0.45,
        "text_contrast": 0.35,
        "dense_text": 0.2,
        "no_text": 0.2,
    }.get(category_code, 0.0)
    return base + bonus



def _teacher_priority_sort_key(issue: dict) -> Tuple[float, int, str]:
    category_code = normalize_whitespace(str(issue.get("category_code") or "")).lower()
    category_rank = {
        "cvd_risk": 0,
        "figure_contrast": 1,
        "small_fonts": 2,
        "text_contrast": 3,
        "dense_text": 4,
        "no_text": 5,
    }.get(category_code, 6)
    unit_index = int(issue.get("unit_index", 0) or 0)
    return (-_teacher_severity_weight(issue), category_rank, f"{unit_index:04d}:{issue.get('location', '')}")



def _teacher_issue_message(issue: dict) -> str:
    category_code = normalize_whitespace(str(issue.get("category_code") or "")).lower()
    mapping = {
        "cvd_risk": "Parte de la información puede depender demasiado del color para interpretarse con claridad.",
        "figure_contrast": "Hay figuras o gráficos cuyo contraste visual puede quedarse corto en proyección o pantalla compartida.",
        "text_contrast": "Hay algunos textos o etiquetas con legibilidad limitada frente al fondo.",
        "small_fonts": "Hay elementos de texto que pueden resultar pequeños para una lectura cómoda en clase.",
        "dense_text": "Hay secciones con demasiada información para una explicación o lectura ágil.",
        "no_text": "Hay una parte del material que conviene revisar manualmente porque no se ha podido interpretar con suficiente fiabilidad.",
    }
    return mapping.get(category_code, normalize_whitespace(str(issue.get("message") or "")) or "Se ha detectado una incidencia visual relevante.")



def _build_teacher_area_stats(issues: List[dict]) -> Dict[str, dict]:
    stats = {
        "color": {"weight": 0.0, "high": 0, "medium": 0, "low": 0, "count": 0},
        "legibility": {"weight": 0.0, "high": 0, "medium": 0, "low": 0, "count": 0},
        "text_load": {"weight": 0.0, "high": 0, "medium": 0, "low": 0, "count": 0},
        "content": {"weight": 0.0, "high": 0, "medium": 0, "low": 0, "count": 0},
        "other": {"weight": 0.0, "high": 0, "medium": 0, "low": 0, "count": 0},
    }
    for issue in issues:
        area_key = _teacher_area_key(issue.get("category_code"))
        severity_code = normalize_whitespace(str(issue.get("severity_code") or "")).lower()
        bucket = stats[area_key]
        bucket["count"] += 1
        bucket["weight"] += _teacher_severity_weight(issue)
        if severity_code in {"high", "medium", "low"}:
            bucket[severity_code] += 1
    return stats



def _teacher_color_signal(payload: dict, issues: List[dict]) -> dict:
    visual_metrics = payload.get("visual_metrics", {}) or {}
    figure_summary = visual_metrics.get("figure_metrics_summary", {}) or {}
    min_cvd_ratio = figure_summary.get("min_cvd_ratio")
    figure_contrast_issues = int(visual_metrics.get("figure_contrast_issues", 0) or 0)
    cvd_issue_count = int(visual_metrics.get("cvd_risk_issues", 0) or 0)
    color_issues = [issue for issue in issues if _teacher_area_key(issue.get("category_code")) == "color"]
    high_color = sum(1 for issue in color_issues if normalize_whitespace(str(issue.get("severity_code") or "")).lower() == "high")
    medium_color = sum(1 for issue in color_issues if normalize_whitespace(str(issue.get("severity_code") or "")).lower() == "medium")

    if cvd_issue_count > 0 or high_color > 0:
        level = "high"
    elif figure_contrast_issues > 0 or medium_color > 0:
        level = "medium"
    elif min_cvd_ratio is not None and float(min_cvd_ratio) < 0.90:
        level = "watch"
    else:
        level = "low"

    return {
        "level": level,
        "min_cvd_ratio": min_cvd_ratio,
        "cvd_issue_count": cvd_issue_count,
        "figure_contrast_issues": figure_contrast_issues,
    }



def _teacher_focus_text(area_key: str, issues: List[dict], payload: Optional[dict] = None) -> str:
    payload = payload or {}
    color_signal = _teacher_color_signal(payload, issues)
    visual_metrics = payload.get("visual_metrics", {}) or {}
    text_contrast_issues = int(visual_metrics.get("text_contrast_issues", 0) or 0)
    small_font_hits = int(visual_metrics.get("small_font_hits", 0) or 0)

    if color_signal["level"] == "high":
        return (
            "La principal mejora pendiente está en cómo se distingue la información visual en algunas figuras y esquemas. "
            "Parte del material puede depender demasiado del color para transmitir significado."
        )
    if color_signal["level"] == "medium":
        return (
            "Conviene reforzar la forma de distinguir la información visual en figuras y gráficos. "
            "En algunas zonas, el contraste o la diferenciación por color puede quedarse corto para clase o proyección."
        )
    if color_signal["level"] == "watch":
        return "La legibilidad general es adecuada, pero conviene revisar algunas figuras donde la diferenciación visual puede depender demasiado del color."
    if area_key == "legibility":
        if small_font_hits and text_contrast_issues:
            return "La principal mejora pendiente está en el tamaño de letra y en algunos elementos de texto con legibilidad limitada."
        if small_font_hits:
            return "La principal mejora pendiente está en el tamaño de letra para proyección o pantalla compartida."
        return "La principal mejora pendiente está en reforzar la legibilidad del texto frente al fondo."
    if area_key == "text_load":
        return "La principal mejora pendiente está en simplificar la cantidad de información por página o diapositiva."
    if area_key == "content":
        return "Conviene revisar manualmente una parte del material porque no se ha podido interpretar con suficiente fiabilidad."
    return "El material presenta un resultado general correcto, con mejoras puntuales recomendables antes de usarlo en clase."



def _teacher_decision_block(payload: dict, issues: List[dict]) -> dict:
    summary = payload.get("summary", {}) or {}
    totals = summary.get("totals", {}) or {}
    severity_totals = summary.get("severity_totals", {}) or {}
    visual_metrics = payload.get("visual_metrics", {}) or {}
    reportable_count = int(totals.get("reportable_issue_count", totals.get("issue_count", 0)) or 0)
    high_count = int(severity_totals.get("alta", 0) or 0)
    medium_count = int(severity_totals.get("media", 0) or 0)
    text_contrast_issues = int(visual_metrics.get("text_contrast_issues", 0) or 0)
    small_font_hits = int(visual_metrics.get("small_font_hits", 0) or 0)

    stats = _build_teacher_area_stats(issues)
    color_signal = _teacher_color_signal(payload, issues)
    if color_signal["level"] == "high":
        stats["color"]["weight"] += 6.0
    elif color_signal["level"] == "medium":
        stats["color"]["weight"] += 3.5
    elif color_signal["level"] == "watch":
        stats["color"]["weight"] += 2.0
    if text_contrast_issues > 0:
        stats["legibility"]["weight"] += min(2.0, 0.35 * text_contrast_issues)
    if small_font_hits > 0:
        stats["legibility"]["weight"] += min(2.0, 0.25 * small_font_hits)

    dominant_area = max(stats.items(), key=lambda item: (item[1]["weight"], item[1]["high"], item[1]["medium"], item[1]["count"]))[0] if issues or any(v.get("weight") for v in stats.values()) else "other"
    focus_text = _teacher_focus_text(dominant_area, issues, payload=payload)

    if reportable_count == 0:
        label = "Muy bien"
        decision = "Se puede usar tal como está."
        decision_summary = "El análisis automático no ha detectado incidencias visuales prioritarias en el material."
    elif high_count >= 2 or reportable_count >= 8:
        label = "Revisar"
        decision = "Conviene revisarlo antes de usarlo en clase."
        decision_summary = f"Se han detectado {reportable_count} incidencias prioritarias. La revisión previa puede mejorar la legibilidad, la interpretación visual o ambas."
    elif high_count >= 1 or medium_count >= 3 or reportable_count >= 4 or color_signal["level"] == "high":
        label = "Atención"
        decision = "Se puede usar, pero conviene corregir algunos puntos antes de clase."
        decision_summary = f"Se han detectado {reportable_count} incidencias que conviene revisar para reforzar la claridad visual del material."
    else:
        label = "Bien"
        decision = "Se puede usar con ajustes menores."
        decision_summary = "El material presenta una base adecuada, aunque conviene revisar algunos detalles antes de proyectarlo o compartirlo."

    return {
        "label": label,
        "decision": decision,
        "summary": decision_summary,
        "focus": focus_text,
    }



def _build_teacher_priorities(payload: dict, issues: List[dict], max_items: int = 3) -> List[dict]:
    priorities = []
    asset_map = {item.get("issue_id"): item for item in (payload.get("teacher_priority_assets", []) or []) if item.get("issue_id")}
    for issue in sorted(issues, key=_teacher_priority_sort_key)[:max_items]:
        category_code = normalize_whitespace(str(issue.get("category_code") or "")).lower()
        action = normalize_whitespace(str(issue.get("recommendation") or "")) or "Revisa esta parte del material antes de usarlo."
        if category_code in {"cvd_risk", "figure_contrast"}:
            action = "Añade etiquetas directas, iconos, patrones o diferencias de forma para que la información no dependa solo del color."
        elif category_code == "small_fonts":
            action = "Aumenta el tamaño de letra y comprueba la legibilidad en proyección o pantalla compartida."
        elif category_code == "text_contrast":
            action = "Refuerza el contraste entre texto y fondo, especialmente en etiquetas, notas breves y elementos secundarios."
        elif category_code == "dense_text":
            action = "Divide la información en bloques más breves y reparte mejor el contenido entre diapositivas o páginas."
        asset = asset_map.get(issue.get("id"), {})
        priorities.append({
            "issue_id": issue.get("id"),
            "location": issue.get("location", "Documento"),
            "area": _teacher_area_label(str(issue.get("category_code") or "")),
            "severity": issue.get("severity", "-"),
            "what": _teacher_issue_message(issue),
            "action": action,
            "preview_image_path": asset.get("preview_image_path"),
            "preview_alt": asset.get("preview_alt"),
            "preview_caption": asset.get("preview_caption"),
        })
    return priorities



def _build_teacher_categories(payload: dict, issues: List[dict]) -> List[dict]:
    color_signal = _teacher_color_signal(payload, issues)
    buckets = [
        {"key": "color", "title": "Color, daltonismo y figuras", "codes": {"cvd_risk", "figure_contrast"}},
        {"key": "legibility", "title": "Tamaño de letra y legibilidad", "codes": {"text_contrast", "small_fonts"}},
        {"key": "text_load", "title": "Cantidad de texto", "codes": {"dense_text"}},
    ]
    categories = []
    for bucket in buckets:
        bucket_issues = [issue for issue in issues if normalize_whitespace(str(issue.get("category_code") or "")).lower() in bucket["codes"]]
        high_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("severity_code") or "")).lower() == "high")
        medium_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("severity_code") or "")).lower() == "medium")
        if bucket["key"] == "color" and not bucket_issues and color_signal["level"] == "watch":
            status = "Bien"
        elif not bucket_issues:
            status = "Muy bien"
        elif high_count >= 1:
            status = "Revisar"
        elif medium_count >= 1 or len(bucket_issues) >= 2:
            status = "Atención"
        else:
            status = "Bien"

        if bucket["key"] == "color":
            cvd_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("category_code") or "")).lower() == "cvd_risk")
            figure_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("category_code") or "")).lower() == "figure_contrast")
            if cvd_count and figure_count:
                detail = "Conviene revisar algunas figuras porque combinan dependencia del color y contraste visual insuficiente."
            elif cvd_count:
                detail = "Hay figuras o esquemas donde la interpretación puede depender demasiado del color."
            elif figure_count:
                detail = "Hay figuras con contraste visual insuficiente que conviene reforzar antes de clase."
            elif color_signal["level"] == "watch":
                detail = "No aparecen incidencias duras en esta categoría, pero sí señales de dependencia cromática que conviene revisar."
            else:
                detail = "No se han detectado incidencias prioritarias en color, daltonismo o figuras."
        elif bucket["key"] == "legibility":
            text_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("category_code") or "")).lower() == "text_contrast")
            font_count = sum(1 for issue in bucket_issues if normalize_whitespace(str(issue.get("category_code") or "")).lower() == "small_fonts")
            if text_count and font_count:
                detail = "Hay elementos de texto que conviene revisar por contraste o por tamaño de letra."
            elif font_count:
                detail = "Se observan textos potencialmente pequeños para una lectura cómoda en proyección."
            elif text_count:
                detail = "Hay algunos textos o etiquetas con legibilidad limitada frente al fondo."
            else:
                detail = "La legibilidad general del texto es adecuada."
        else:
            if bucket_issues:
                detail = "Hay secciones con demasiada información para una explicación o lectura ágil."
            else:
                detail = "La carga textual general es adecuada y no se aprecia sobrecarga relevante."
        categories.append({"title": bucket["title"], "status": status, "count": len(bucket_issues), "detail": detail})
    return categories



def _build_teacher_strengths(payload: dict, categories: List[dict]) -> List[str]:
    strengths: List[str] = []
    visual_metrics = payload.get("visual_metrics", {}) or {}
    if int(visual_metrics.get("dense_text_sections", 0) or 0) == 0:
        strengths.append("La cantidad general de texto es adecuada para un seguimiento razonable en clase.")
    for category in categories:
        if normalize_whitespace(str(category.get("status") or "")) == "Muy bien":
            detail = normalize_whitespace(str(category.get("detail") or ""))
            title = normalize_whitespace(str(category.get("title") or ""))
            if detail:
                strengths.append(f"{title}: {detail}")
    deduped = []
    for item in strengths:
        if item not in deduped:
            deduped.append(item)
    return deduped[:3]




def _set_mupdf_warning_display(enabled: bool) -> Optional[bool]:
    try:
        import fitz  # type: ignore

        tools = getattr(fitz, "TOOLS", None)
        if tools is None or not hasattr(tools, "mupdf_display_warnings"):
            return None
        previous = tools.mupdf_display_warnings()
        try:
            if hasattr(tools, "reset_mupdf_warnings"):
                tools.reset_mupdf_warnings()
        except Exception:
            pass
        tools.mupdf_display_warnings(bool(enabled))
        return bool(previous)
    except Exception:
        return None



def _restore_mupdf_warning_display(previous: Optional[bool]) -> None:
    if previous is None:
        return
    try:
        import fitz  # type: ignore

        tools = getattr(fitz, "TOOLS", None)
        if tools is None or not hasattr(tools, "mupdf_display_warnings"):
            return
        tools.mupdf_display_warnings(bool(previous))
        try:
            if hasattr(tools, "reset_mupdf_warnings"):
                tools.reset_mupdf_warnings()
        except Exception:
            pass
    except Exception:
        pass

def _prepare_preview_pdf(input_file: Path, logger: logging.Logger) -> Tuple[Optional[Path], Optional[Path]]:
    suffix = input_file.suffix.lower()
    if suffix == ".pdf":
        return input_file, None

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        logger.warning("No se pudo generar previsualización de prioridades: LibreOffice no está disponible.")
        return None, None

    tmp_dir = Path(tempfile.mkdtemp(prefix="aiuda_preview_pdf_"))
    cmd = [
        soffice,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(tmp_dir),
        str(input_file),
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        logger.warning("No se pudo generar PDF de previsualización para prioridades: %s", proc.stderr.strip() or proc.stdout.strip())
        shutil.rmtree(tmp_dir, ignore_errors=True)
        return None, None

    candidates = list(tmp_dir.glob("*.pdf"))
    if not candidates:
        logger.warning("No se generó ningún PDF de previsualización para prioridades.")
        shutil.rmtree(tmp_dir, ignore_errors=True)
        return None, None
    return candidates[0], tmp_dir


def _clip_rect_with_padding(rect: Any, bbox: Optional[List[float]]) -> Any:
    if bbox is None:
        return rect
    try:
        import fitz  # type: ignore

        clip = fitz.Rect(*bbox)
        if clip.is_empty or clip.is_infinite:
            return rect
        pad_x = max((clip.x1 - clip.x0) * 0.14, rect.width * 0.02)
        pad_y = max((clip.y1 - clip.y0) * 0.18, rect.height * 0.02)
        clip = fitz.Rect(clip.x0 - pad_x, clip.y0 - pad_y, clip.x1 + pad_x, clip.y1 + pad_y)
        min_w = rect.width * 0.18
        min_h = rect.height * 0.16
        if clip.width < min_w:
            extra = (min_w - clip.width) / 2.0
            clip = fitz.Rect(clip.x0 - extra, clip.y0, clip.x1 + extra, clip.y1)
        if clip.height < min_h:
            extra = (min_h - clip.height) / 2.0
            clip = fitz.Rect(clip.x0, clip.y0 - extra, clip.x1, clip.y1 + extra)
        clip = clip & rect
        return clip if not clip.is_empty else rect
    except Exception:
        return rect


def _save_priority_preview(page: Any, bbox: Optional[List[float]], output_path: Path) -> bool:
    try:
        import fitz  # type: ignore

        clip = _clip_rect_with_padding(page.rect, bbox)
        zoom = 2.0 if bbox else 1.25
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=clip, alpha=False)
        pix.save(str(output_path))
        return True
    except Exception:
        return False


def _build_teacher_priority_assets(input_file: Path, output_dir: Path, payload: dict, logger: logging.Logger, max_items: int = 3) -> List[dict]:
    issues = _filter_reportable_issues(payload)
    if not issues:
        return []

    ranked = sorted(issues, key=_teacher_priority_sort_key)[:max_items]
    preview_pdf, cleanup_dir = _prepare_preview_pdf(input_file, logger)
    if preview_pdf is None:
        return []

    asset_dir = output_dir / "_teacher_assets"
    ensure_dir(asset_dir)
    assets: List[dict] = []
    previous_warning_state = _set_mupdf_warning_display(False)

    try:
        import fitz  # type: ignore

        doc = fitz.open(str(preview_pdf))
        try:
            for idx, issue in enumerate(ranked, start=1):
                if len(doc) == 0:
                    break
                page_index = max(0, min(len(doc) - 1, int(issue.get("unit_index", 1) or 1) - 1))
                page = doc[page_index]
                bbox = normalize_bbox(issue.get("bbox"))
                rel_path = f"_teacher_assets/priority_{idx:02d}.png"
                out_path = output_dir / rel_path
                if _save_priority_preview(page, bbox, out_path):
                    assets.append({
                        "issue_id": issue.get("id"),
                        "preview_image_path": rel_path,
                        "preview_alt": f"Vista de apoyo de {issue.get('location', 'la incidencia prioritaria')}",
                        "preview_caption": issue.get("bbox_human") or issue.get("location") or "Vista de apoyo",
                        "preview_source_page": page_index + 1,
                    })
        finally:
            doc.close()
    except Exception as exc:
        logger.warning("No se pudieron generar imágenes de apoyo para las prioridades docentes: %s", exc)
    finally:
        _restore_mupdf_warning_display(previous_warning_state)
        if cleanup_dir is not None:
            shutil.rmtree(cleanup_dir, ignore_errors=True)

    return assets


def render_teacher_report_text(payload: dict) -> str:
    issues = _filter_reportable_issues(payload)
    recommendations = clean_list_items(payload.get("recommendations", []) or [])
    decision = _teacher_decision_block(payload, issues)
    priorities = _build_teacher_priorities(payload, issues, max_items=3)
    categories = _build_teacher_categories(payload, issues)
    strengths = _build_teacher_strengths(payload, categories)

    lines = [
        "Informe para profesorado",
        "",
        f"Resultado para docencia: {decision.get('label', '-')}",
        f"Decisión rápida: {decision.get('decision', '-')}",
        f"Resumen ejecutivo: {decision.get('summary', '-')}",
        f"Foco principal: {decision.get('focus', '-')}",
        "",
        "Qué conviene corregir primero",
    ]
    if priorities:
        for item in priorities:
            lines.extend([
                f"- {item.get('location', '-')}",
                f"  Área: {item.get('area', '-')}",
                f"  Qué ocurre: {item.get('what', '-')}",
                f"  Qué hacer: {item.get('action', '-')}",
                "",
            ])
    else:
        lines.extend(["- No se detectaron incidencias prioritarias en este análisis automático.", ""])

    lines.append("Resumen rápido por categorías")
    for category in categories:
        lines.append(f"- {category.get('title', '-')}: {category.get('status', '-')}")
        lines.append(f"  {category.get('detail', '-')}")

    lines.extend(["", "Qué ya funciona bien"])
    if strengths:
        for item in strengths:
            lines.append(f"- {item}")
    else:
        lines.append("- El material presenta una base general utilizable, aunque conviene revisar algunos detalles antes de clase.")

    lines.extend(["", "Recomendaciones prácticas"])
    if recommendations:
        for item in recommendations[:5]:
            lines.append(f"- {item}")
    else:
        lines.append("- No hay recomendaciones adicionales.")

    lines.extend(["", "A continuación se incluye un anexo técnico con el detalle completo del análisis automático, para consulta y revisión más especializada.", ""])
    return "\n".join(lines) + "\n"



def render_technical_report_text(payload: dict) -> str:
    summary = payload.get("summary", {}) or {}
    totals = summary.get("totals", {}) or {}
    severity_totals = summary.get("severity_totals", {}) or {}
    all_issues = payload.get("issues", []) or []
    issues = [issue for issue in all_issues if issue.get("category_code") != "processing_note"]
    recommendations = clean_list_items(payload.get("recommendations", []) or [])
    visual_metrics = payload.get("visual_metrics", {}) or {}
    methodology = payload.get("methodology", {}) or {}

    lines = [
        "Anexo técnico",
        "",
        summary.get("title", "Informe de accesibilidad documental"),
        "",
        summary.get("overview", ""),
        "",
        "Resumen",
        f"- Unidades analizadas: {totals.get('units_analyzed', '-')}",
        f"- Alertas visuales: {totals.get('reportable_issue_count', totals.get('issue_count', '-'))}",
        f"- Notas de procesamiento: {totals.get('processing_note_count', 0)}",
        f"- Tipo de documento: {totals.get('document_type', '-')}",
        f"- Bloques textuales evaluados: {totals.get('text_blocks_analyzed', '-')}",
        f"- Figuras analizadas: {totals.get('figures_analyzed', '-')}",
        f"- Severidad alta/media/baja: {severity_totals.get('alta', 0)}/{severity_totals.get('media', 0)}/{severity_totals.get('baja', 0)}",
        "",
        "Incidencias prioritarias",
    ]

    if issues:
        for issue in issues:
            lines.extend([
                f"- {issue.get('location', '-')}: {issue.get('category', '-')}",
                f"  Severidad: {issue.get('severity', '-')}",
                f"  Mensaje: {issue.get('message', '-')}",
                f"  Recomendación: {issue.get('recommendation', '-')}",
                (f"  Extracto: {issue.get('excerpt', '')}" if issue.get("excerpt") else ""),
                (f"  Referencia: {issue.get('reference_code', '')}" if issue.get("reference_code") else ""),
                "",
            ])
    else:
        lines.append("- No se detectaron incidencias visuales prioritarias en este análisis automático.")
        lines.append("")

    lines.extend([
        "Métricas visuales",
        f"- Contrastes textuales evaluados: {visual_metrics.get('text_blocks_analyzed', '-')}",
        f"- Incidencias de contraste textual: {visual_metrics.get('text_contrast_issues', '-')}",
        f"- Secciones con carga visual alta: {visual_metrics.get('dense_text_sections', '-')}",
        f"- Alertas de fuente pequeña: {visual_metrics.get('small_font_hits', '-')}",
        f"- Figuras detectadas: {visual_metrics.get('figure_count', '-')}",
        f"- Incidencias de contraste en figuras: {visual_metrics.get('figure_contrast_issues', '-')}",
        f"- Incidencias de riesgo CVD: {visual_metrics.get('cvd_risk_issues', '-')}",
        "",
    ])

    text_compact_note = build_text_metric_compact_note(visual_metrics)
    figure_compact_note = build_figure_metric_compact_note(visual_metrics)

    lines.append("Resumen técnico de contraste textual")
    if text_compact_note:
        lines.append(f"- {text_compact_note}")
    else:
        for item in build_metric_summary_lines(visual_metrics.get("text_metrics_summary", {}) or {}, "text"):
            lines.append(f"- {item}")
        text_rows = _select_relevant_metric_rows(visual_metrics.get("text_metric_samples", []) or [])
        if text_rows:
            lines.extend(["", "Muestras de contraste textual"])
            for row in text_rows:
                lines.append(f"- {row.get('location', '-')}, bloque {row.get('text_block_index', '-')}: {row.get('status', '-')} · {row.get('ratio', '-')}:1 (umbral {row.get('threshold', '-')}:1) · {row.get('bbox_human', '-')}")

    lines.append("")
    lines.append("Resumen técnico de figuras y CVD")
    if figure_compact_note:
        lines.append(f"- {figure_compact_note}")
    else:
        for item in build_metric_summary_lines(visual_metrics.get("figure_metrics_summary", {}) or {}, "figure"):
            lines.append(f"- {item}")
        figure_rows = _select_relevant_metric_rows(visual_metrics.get("figure_metric_samples", []) or [])
        if figure_rows:
            lines.extend(["", "Muestras de figuras y CVD"])
            for row in figure_rows:
                ratio_text = f"{row.get('ratio')}:1" if row.get('ratio') is not None else "sin dato"
                cvd_text = row.get('worst_cvd_ratio', '-')
                lines.append(f"- {row.get('location', '-')}, figura {row.get('figure_index', '-')}: {row.get('status', '-')} · contraste {ratio_text} · peor score CVD {cvd_text} · {row.get('bbox_human', '-')}")


    processing_notes = clean_list_items(visual_metrics.get("processing_notes", []) or [])
    if processing_notes:
        lines.extend(["", "Notas de procesamiento"])
        for note in processing_notes:
            lines.append(f"- {note}")

    lines.extend(["", "Recomendaciones generales"])
    for item in recommendations:
        lines.append(f"- {item}")

    lines.extend(["", methodology.get("title", "Metodología y referencias")])
    for rule in methodology.get("contrast_rules", []) or []:
        threshold = rule.get("threshold")
        threshold_text = f" · Umbral: {threshold}:1" if threshold is not None else ""
        lines.append(f"- {rule.get('label', '-')}: {rule.get('metric_type', '-')} · {rule.get('reference_code', '-')}{threshold_text}")
        if rule.get("details"):
            lines.append(f"  {rule.get('details')}")
    if methodology.get("limitations"):
        lines.append("")
        lines.append("Limitaciones de la estimación")
        for item in clean_list_items(methodology.get("limitations", [])):
            lines.append(f"- {item}")
    if methodology.get("references"):
        lines.append("")
        lines.append("Referencias")
        for ref in methodology.get("references", []):
            lines.append(f"- {ref.get('title', '-')}")
            if ref.get("details"):
                lines.append(f"  {ref.get('details')}")

    return "\n".join(lines) + "\n"



def render_report_text(payload: dict) -> str:
    return render_teacher_report_text(payload) + render_technical_report_text(payload)

def build_summary_cards_html(summary: dict, visual_metrics: dict) -> str:
    totals = summary.get("totals", {}) or {}
    severity_totals = summary.get("severity_totals", {}) or {}
    cards = [
        ("Unidades analizadas", totals.get("units_analyzed", "-")),
        ("Alertas visuales", totals.get("reportable_issue_count", totals.get("issue_count", "-"))),
        ("Notas de procesamiento", totals.get("processing_note_count", 0)),
        ("Bloques textuales evaluados", totals.get("text_blocks_analyzed", "-")),
        ("Figuras analizadas", totals.get("figures_analyzed", "-")),
        ("Incidencias de contraste textual", visual_metrics.get("text_contrast_issues", 0)),
        ("Incidencias de riesgo CVD", visual_metrics.get("cvd_risk_issues", 0)),
        ("Severidad alta", severity_totals.get("alta", 0)),
    ]
    return "".join(
        f"<div class='metric-card'><div class='metric-label'>{html.escape(str(label))}</div><div class='metric-value'>{html.escape(str(value))}</div></div>"
        for label, value in cards
    )

def build_metric_table_html(rows: Iterable[dict], columns: List[Tuple[str, str]], empty_text: str = "No hay muestras suficientes para mostrar en esta sección.") -> str:
    rows = [row for row in rows if row]
    if not rows:
        return f"<div class='muted'>{html.escape(empty_text)}</div>"
    head = "".join(f"<th>{html.escape(label)}</th>" for _, label in columns)
    body_parts = []
    for row in rows:
        row_class = ""
        status = str(row.get("status", "")).lower()
        if "incidencia" in status or "riesgo" in status:
            row_class = " class='row-alert'"
        elif "cercano" in status:
            row_class = " class='row-warn'"
        cols = []
        for key, _label in columns:
            cols.append(f"<td>{html.escape(render_metric_value(row.get(key)))}</td>")
        body_parts.append(f"<tr{row_class}>" + "".join(cols) + "</tr>")
    return f"<table class='data-table'><thead><tr>{head}</tr></thead><tbody>{''.join(body_parts)}</tbody></table>"


def build_metric_summary_list_html(lines: List[str]) -> str:
    lines = clean_list_items(lines)
    if not lines:
        return "<div class='muted'>No hay datos agregados disponibles.</div>"
    return "<ul class='summary-list'>" + "".join(f"<li>{html.escape(item)}</li>" for item in lines) + "</ul>"

def _metric_row_is_relevant(row: dict) -> bool:
    status = normalize_whitespace(str((row or {}).get("status", ""))).lower()
    return any(token in status for token in ("incidencia", "riesgo", "cercano"))


def _select_relevant_metric_rows(rows: Iterable[dict]) -> List[dict]:
    selected = [row for row in (rows or []) if row and _metric_row_is_relevant(row)]
    return selected


def build_text_metric_compact_note(visual_metrics: dict) -> Optional[str]:
    summary = visual_metrics.get("text_metrics_summary", {}) or {}
    total = int(visual_metrics.get("text_blocks_total", 0) or 0)
    analyzed = int(visual_metrics.get("text_blocks_analyzed", 0) or 0)
    below = int(summary.get("below_threshold_count", 0) or 0)
    near = int(summary.get("near_threshold_count", 0) or 0)

    if analyzed <= 0:
        if total <= 0:
            return "No se obtuvieron métricas estructuradas de contraste textual aplicables a este documento."
        return "No se generaron métricas estructuradas de contraste textual suficientemente fiables como para mostrar una tabla detallada en esta sección."

    if below == 0 and near == 0:
        return f"Se evaluaron {analyzed} bloques de texto y no se detectaron incidencias prioritarias de contraste textual."

    return None


def build_figure_metric_compact_note(visual_metrics: dict) -> Optional[str]:
    figure_count = int(visual_metrics.get("figure_count", 0) or 0)
    analyzed = int(visual_metrics.get("figures_analyzed", 0) or 0)
    contrast_issues = int(visual_metrics.get("figure_contrast_issues", 0) or 0)
    cvd_issues = int(visual_metrics.get("cvd_risk_issues", 0) or 0)
    skip_reasons = visual_metrics.get("figure_skip_reasons", {}) or {}

    if contrast_issues > 0 or cvd_issues > 0:
        return None

    if figure_count <= 0:
        return "No se han detectado figuras relevantes en el documento, por lo que no se muestra un bloque técnico ampliado de figuras y CVD."

    if analyzed > 0:
        return f"Se detectaron {figure_count} figuras y no presentan incidencias prioritarias de contraste o riesgo CVD en este análisis automático."

    if skip_reasons:
        return "Se detectaron elementos visuales, pero no requieren aquí un desarrollo técnico ampliado; en este contexto parecen principalmente ilustrativos o no críticos para el análisis."

    return "Se detectaron figuras, pero no presentan incidencias prioritarias y por ello esta sección se muestra de forma compacta."



def build_teacher_summary_cards_html(categories: List[dict]) -> str:
    if not categories:
        return "<div class='muted'>No hay datos suficientes para mostrar un resumen docente.</div>"
    return "".join(
        f"<div class='metric-card teacher-card'><div class='metric-label'>{html.escape(str(item.get('title', '-')))}</div><div class='metric-value metric-value-small'>{html.escape(str(item.get('status', '-')))}</div><div class='teacher-detail'>{html.escape(str(item.get('detail', '-')))}</div></div>"
        for item in categories
    )



def build_teacher_priorities_html(priorities: List[dict]) -> str:
    if not priorities:
        return "<div class='muted'>No se detectaron incidencias prioritarias en este análisis automático.</div>"
    cards = []
    for item in priorities:
        image_html = ""
        if item.get("preview_image_path"):
            caption = html.escape(str(item.get("preview_caption") or "Vista de apoyo"))
            alt = html.escape(str(item.get("preview_alt") or "Vista de apoyo"))
            src = html.escape(str(item.get("preview_image_path")))
            image_html = f"""
              <figure class="priority-figure">
                <img src="{src}" alt="{alt}">
                <figcaption class="priority-caption">{caption}</figcaption>
              </figure>
            """
        cards.append(
            f"""
            <article class="card issue-card teacher-priority-card">
              <div class="teacher-priority-layout">
                <div class="teacher-priority-text">
                  <div class="issue-location">{html.escape(str(item.get('location', '-')))}</div>
                  <div class="issue-meta">{html.escape(str(item.get('area', '-')))} · {html.escape(str(item.get('severity', '-')))}</div>
                  <div class="issue-message"><strong>Qué ocurre:</strong> {html.escape(str(item.get('what', '-')))}</div>
                  <div class="issue-recommendation"><strong>Qué hacer:</strong> {html.escape(str(item.get('action', '-')))}</div>
                </div>
                {image_html}
              </div>
            </article>
            """
        )
    return "".join(cards)

def render_report_html(payload: dict, title: str) -> str:
    summary = payload.get("summary", {}) or {}
    issues = _filter_reportable_issues(payload)
    recommendations = clean_list_items(payload.get("recommendations", []) or [])
    visual_metrics = payload.get("visual_metrics", {}) or {}
    methodology = payload.get("methodology", {}) or {}

    decision = _teacher_decision_block(payload, issues)
    priorities = _build_teacher_priorities(payload, issues, max_items=3)
    categories = _build_teacher_categories(payload, issues)
    strengths = _build_teacher_strengths(payload, categories)

    if issues:
        cards = []
        for issue in issues:
            meta_parts = [issue.get("category", "-"), issue.get("severity", "-")]
            if issue.get("reference_code"):
                meta_parts.append(issue.get("reference_code"))
            if issue.get("metric_type"):
                meta_parts.append(issue.get("metric_type"))
            excerpt_html = ""
            if issue.get("excerpt"):
                excerpt_html = f"<div class='excerpt'><strong>Extracto:</strong> {html.escape(str(issue['excerpt']))}</div>"
            metric_html = ""
            metric_bits = []
            if issue.get("ratio") is not None and issue.get("threshold") is not None:
                metric_bits.append(f"Ratio: {html.escape(str(issue.get('ratio')))}:1")
                metric_bits.append(f"Umbral: {html.escape(str(issue.get('threshold')))}:1")
            if issue.get("score_ratio") is not None:
                metric_bits.append(f"Score CVD: {html.escape(str(issue.get('score_ratio')))}")
            if issue.get("bbox_human"):
                metric_bits.append(f"Zona: {html.escape(str(issue.get('bbox_human')))}")
            if issue.get("color_fg") and issue.get("color_bg"):
                metric_bits.append(f"Colores: {html.escape(str(issue.get('color_fg')))} sobre {html.escape(str(issue.get('color_bg')))}")
            if metric_bits:
                metric_html = "<div class='issue-metrics'>" + " · ".join(metric_bits) + "</div>"
            cards.append(
                f"""
                <article class="card issue-card">
                  <div class="issue-location">{html.escape(str(issue.get('location', '-')))}</div>
                  <div class="issue-meta">{' · '.join(html.escape(str(x)) for x in meta_parts if x)}</div>
                  <div class="issue-message">{html.escape(str(issue.get('message', '-')))}</div>
                  {metric_html}
                  <div class="issue-recommendation"><strong>Recomendación:</strong> {html.escape(str(issue.get('recommendation', '-')))}</div>
                  {excerpt_html}
                </article>
                """
            )
        issues_html = "".join(cards)
    else:
        issues_html = "<div class='muted'>No se detectaron incidencias visuales prioritarias en este análisis automático.</div>"

    teacher_recommendations_html = "".join(f"<li>{html.escape(str(item))}</li>" for item in recommendations[:5]) or "<li>No hay recomendaciones adicionales.</li>"
    recommendations_html = "".join(f"<li>{html.escape(str(item))}</li>" for item in recommendations) or "<li>No hay recomendaciones adicionales.</li>"

    text_compact_note = build_text_metric_compact_note(visual_metrics)
    figure_compact_note = build_figure_metric_compact_note(visual_metrics)
    text_metric_rows = _select_relevant_metric_rows(visual_metrics.get("text_metric_samples", []) or [])
    figure_metric_rows = _select_relevant_metric_rows(visual_metrics.get("figure_metric_samples", []) or [])

    if text_compact_note:
        text_summary_html = ""
        text_metrics_html = f"<div class='muted compact-note'>{html.escape(text_compact_note)}</div>"
    else:
        text_summary_html = build_metric_summary_list_html(build_metric_summary_lines(visual_metrics.get("text_metrics_summary", {}) or {}, "text"))
        text_metrics_html = build_metric_table_html(
            text_metric_rows,
            columns=[
                ("location", "Ubicación"),
                ("text_block_index", "Bloque"),
                ("status", "Estado"),
                ("ratio", "Ratio"),
                ("threshold", "Umbral"),
                ("bbox_human", "Zona"),
                ("method", "Método"),
            ],
            empty_text="No hay muestras relevantes de contraste textual para mostrar.",
        )

    if figure_compact_note:
        figure_summary_html = ""
        figure_metrics_html = f"<div class='muted compact-note'>{html.escape(figure_compact_note)}</div>"
    else:
        figure_summary_html = build_metric_summary_list_html(build_metric_summary_lines(visual_metrics.get("figure_metrics_summary", {}) or {}, "figure"))
        figure_metrics_html = build_metric_table_html(
            figure_metric_rows,
            columns=[
                ("location", "Ubicación"),
                ("figure_index", "Figura"),
                ("status", "Estado"),
                ("ratio", "Ratio figura"),
                ("worst_cvd_type", "CVD más sensible"),
                ("worst_cvd_ratio", "Score CVD"),
                ("bbox_human", "Zona"),
            ],
            empty_text="No hay muestras relevantes de figuras para mostrar.",
        )

    processing_notes = clean_list_items(visual_metrics.get("processing_notes", []) or [])
    processing_notes_html = "".join(f"<li>{html.escape(str(note))}</li>" for note in processing_notes) or "<li>No hay notas adicionales.</li>"

    reference_items = "".join(
        f"<li><strong>{html.escape(str(ref.get('metric_type', '-')))}</strong> · {html.escape(str(ref.get('title', '-')))}"
        + (f"<div class='reference-detail'>{html.escape(str(ref.get('details', '')))}</div>" if ref.get("details") else "")
        + "</li>"
        for ref in (methodology.get("references", []) or [])
    )
    contrast_rules_html = "".join(
        f"<li><strong>{html.escape(str(rule.get('label', '-')))}</strong> · {html.escape(str(rule.get('metric_type', '-')))}"
        + (f" · <span class='rule-badge'>Umbral: {html.escape(str(rule.get('threshold')))}:1</span>" if rule.get("threshold") is not None else "")
        + (f" · <span class='rule-badge'>{html.escape(str(rule.get('reference_code', '-')))}</span>" if rule.get("reference_code") else "")
        + (f"<div class='reference-detail'>{html.escape(str(rule.get('details', '')))}</div>" if rule.get("details") else "")
        + "</li>"
        for rule in (methodology.get("contrast_rules", []) or [])
    )
    limitations_html = "".join(f"<li>{html.escape(str(item))}</li>" for item in clean_list_items(methodology.get("limitations", []) or []))

    styles = """
    :root {
      --bg: #f4f7fa;
      --panel: #ffffff;
      --line: #dde6ee;
      --line-strong: #c9d6e1;
      --text: #12263a;
      --muted: #536474;
      --soft: #f8fafc;
      --accent: #184f75;
      --accent-soft: #eef6fb;
      --alert-soft: #fff5f5;
      --warn-soft: #fffaf0;
      --teacher-soft: #f4f8fb;
    }
    * { box-sizing: border-box; }
    body { margin: 0; padding: 24px; background: var(--bg); color: var(--text); font-family: Arial, Helvetica, sans-serif; }
    .page { max-width: 980px; margin: 0 auto; background: var(--panel); border: 1px solid var(--line); border-radius: 22px; overflow: hidden; box-shadow: 0 6px 24px rgba(18, 38, 58, 0.08); }
    .hero { padding: 30px 32px 18px 32px; border-bottom: 1px solid var(--line); background: linear-gradient(180deg, #ffffff 0%, #fbfdff 100%); }
    .hero h1 { margin: 0; font-size: 30px; line-height: 1.2; }
    .hero p { margin: 12px 0 0 0; color: var(--muted); font-size: 15px; line-height: 1.7; }
    .section { padding: 24px 32px 0 32px; }
    .section:last-child { padding-bottom: 32px; }
    h2 { margin: 0 0 14px 0; font-size: 21px; line-height: 1.3; }
    h3 { margin: 0 0 12px 0; font-size: 17px; line-height: 1.3; }
    .section-kicker { color: var(--muted); font-size: 12px; letter-spacing: 0.08em; text-transform: uppercase; margin-bottom: 8px; font-weight: 700; }
    .metric-grid { display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 12px; }
    .metric-card { border: 1px solid var(--line); border-radius: 16px; background: var(--soft); padding: 14px 16px; min-height: 104px; }
    .teacher-card { background: var(--teacher-soft); }
    .metric-label { color: var(--muted); font-size: 13px; line-height: 1.45; margin-bottom: 8px; }
    .metric-value { font-size: 26px; font-weight: 700; color: var(--text); }
    .metric-value-small { font-size: 22px; }
    .teacher-detail { color: var(--muted); font-size: 13px; line-height: 1.6; margin-top: 8px; }
    .card { border: 1px solid var(--line); border-radius: 16px; background: #fff; padding: 16px 18px; margin-bottom: 14px; }
    .decision-card { background: linear-gradient(180deg, #ffffff 0%, #f8fcff 100%); }
    .executive-lead { font-size: 15px; line-height: 1.8; }
    .teacher-note { color: var(--muted); font-size: 14px; line-height: 1.7; }
    .issue-card { break-inside: avoid; page-break-inside: avoid; }
    .teacher-priority-card { background: #fffdf8; }
    .teacher-priority-layout { display: grid; grid-template-columns: minmax(0, 1fr) 220px; gap: 16px; align-items: start; }
    .teacher-priority-text { min-width: 0; }
    .priority-figure { margin: 0; }
    .priority-figure img { width: 100%; display: block; border: 1px solid var(--line-strong); border-radius: 12px; background: #fff; }
    .priority-caption { margin-top: 6px; color: var(--muted); font-size: 12px; line-height: 1.5; }
    .issue-location { font-weight: 700; font-size: 15px; margin-bottom: 6px; }
    .issue-meta { color: var(--muted); font-size: 13px; line-height: 1.5; margin-bottom: 10px; }
    .issue-message, .issue-recommendation, .reference-detail, li, td, th, .muted { font-size: 14px; line-height: 1.7; }
    .issue-metrics { margin: 10px 0; padding: 10px 12px; border-radius: 12px; background: var(--accent-soft); color: var(--accent); font-size: 13px; line-height: 1.6; }
    .excerpt { margin-top: 10px; padding: 10px 12px; background: var(--soft); border-radius: 12px; color: var(--muted); font-size: 13px; line-height: 1.6; }
    .muted { color: var(--muted); }
    .compact-note { padding: 2px 0 0 0; }
    .data-table { width: 100%; border-collapse: collapse; border: 1px solid var(--line); border-radius: 14px; overflow: hidden; background: #fff; }
    .data-table th, .data-table td { border-bottom: 1px solid var(--line); padding: 10px 12px; text-align: left; vertical-align: top; }
    .data-table thead th { background: var(--soft); font-size: 13px; color: var(--muted); font-weight: 700; }
    .row-alert td { background: var(--alert-soft); }
    .row-warn td { background: var(--warn-soft); }
    ul { margin: 0; padding-left: 20px; }
    .summary-list { margin-top: 0; }
    .rule-badge { display: inline-block; padding: 2px 8px; border-radius: 999px; border: 1px solid var(--line-strong); background: var(--soft); font-size: 12px; color: var(--muted); }
    .columns { display: grid; grid-template-columns: 1fr 1fr; gap: 18px; align-items: start; }
    .annex-intro { background: #fbfcfd; }
    .section-break { padding-top: 28px; }
    @media print {
      @page { size: A4; margin: 14mm; }
      body { padding: 0; background: #fff; }
      .page { max-width: none; border: 0; box-shadow: none; border-radius: 0; }
      .hero, .metric-card, .card, .issue-card { break-inside: avoid; page-break-inside: avoid; }
      .columns { display: block; }
      .columns > * { margin-bottom: 14px; }
      .section { padding-left: 0; padding-right: 0; }
      .hero { padding-left: 0; padding-right: 0; }
      .section-break { page-break-before: always; break-before: page; }
    }
    @media (max-width: 900px) { .metric-grid, .columns { grid-template-columns: 1fr 1fr; } }
    @media (max-width: 640px) { body { padding: 12px; } .page { border-radius: 16px; } .hero, .section { padding-left: 18px; padding-right: 18px; } .metric-grid, .columns, .teacher-priority-layout { grid-template-columns: 1fr; } }
    """

    return f"""<!DOCTYPE html>
<html lang="es">
  <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>{html.escape(title)}</title>
    <style>{styles}</style>
  </head>
  <body>
    <main class="page">
      <header class="hero">
        <h1>{html.escape(str(summary.get('title', title)))}</h1>
        <p>{html.escape(str(summary.get('overview', '')))}</p>
      </header>

      <section class="section">
        <div class="section-kicker">Resumen ejecutivo</div>
        <h2>Informe para profesorado</h2>
        <div class="card decision-card">
          <h3>Resultado para docencia: {html.escape(str(decision.get('label', '-')))}</h3>
          <p class="executive-lead"><strong>Decisión rápida:</strong> {html.escape(str(decision.get('decision', '-')))}</p>
          <p>{html.escape(str(decision.get('summary', '-')))}</p>
          <p><strong>Foco principal:</strong> {html.escape(str(decision.get('focus', '-')))}</p>
        </div>
      </section>

      <section class="section">
        <h2>Qué conviene corregir primero</h2>
        {build_teacher_priorities_html(priorities)}
      </section>

      <section class="section">
        <h2>Resumen rápido por categorías</h2>
        <div class="metric-grid">{build_teacher_summary_cards_html(categories)}</div>
      </section>

      <section class="section">
        <h2>Qué ya funciona bien</h2>
        <div class="card teacher-note"><ul>{''.join(f'<li>{html.escape(str(item))}</li>' for item in strengths) or '<li>El material presenta una base general utilizable, aunque conviene revisar algunos detalles antes de clase.</li>'}</ul></div>
      </section>

      <section class="section">
        <h2>Recomendaciones prácticas</h2>
        <div class="card"><ul>{teacher_recommendations_html}</ul></div>
        <div class="card teacher-note"><p>A continuación se incluye un anexo técnico con el detalle completo del análisis automático, para consulta y revisión más especializada.</p></div>
      </section>

      <section class="section section-break technical-annex">
        <h2>Anexo técnico</h2>
        <div class="card annex-intro">
          <p>La siguiente sección conserva el detalle técnico del análisis automático para consulta, trazabilidad y revisión más especializada.</p>
        </div>
      </section>

      <section class="section">
        <h2>Incidencias prioritarias</h2>
        {issues_html}
      </section>

      <section class="section">
        <h2>Recomendaciones generales</h2>
        <div class="card"><ul>{recommendations_html}</ul></div>
      </section>

      <section class="section">
        <h2>Métricas de contraste y accesibilidad visual</h2>
        <div class="columns">
          <div class="card">
            <h3>Contraste textual</h3>
            {text_summary_html}
            {text_metrics_html}
          </div>
          <div class="card">
            <h3>Figuras y riesgo CVD</h3>
            {figure_summary_html}
            {figure_metrics_html}
          </div>
        </div>
      </section>

      <section class="section">
        <h2>{html.escape(str(methodology.get('title', 'Metodología y referencias')))}</h2>
        <div class="card">
          <h3>Criterios de referencia</h3>
          <ul>{contrast_rules_html}</ul>
        </div>
        <div class="columns" style="margin-top:18px;">
          <div class="card">
            <h3>Limitaciones de la estimación</h3>
            <ul>{limitations_html or '<li>No se han registrado limitaciones adicionales.</li>'}</ul>
          </div>
          <div class="card">
            <h3>Referencias</h3>
            <ul>{reference_items}</ul>
          </div>
        </div>
      </section>
    </main>
  </body>
</html>
"""

def render_html_to_pdf(html_content: str, pdf_path: Path, payload: Optional[dict] = None) -> None:
    """Genera un PDF a partir del HTML del informe.

    Intenta usar WeasyPrint para conservar el estilo del HTML. Si no está
    disponible, recurre a un fallback estructurado con ReportLab para no romper
    el flujo del sistema y mantener mejor paginación que un volcado lineal.
    """
    try:
        from weasyprint import HTML  # type: ignore

        HTML(string=html_content, base_url=str(pdf_path.parent)).write_pdf(str(pdf_path))
        return
    except Exception:
        pass

    try:
        from reportlab.lib import colors  # type: ignore
        from reportlab.lib.enums import TA_LEFT  # type: ignore
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
        from reportlab.lib.units import mm  # type: ignore
        from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # type: ignore

        payload = payload or {}
        summary = payload.get("summary", {}) or {}
        visual_metrics = payload.get("visual_metrics", {}) or {}
        all_issues = payload.get("issues", []) or []
        issues = [issue for issue in all_issues if issue.get("category_code") != "processing_note"]
        recommendations = clean_list_items(payload.get("recommendations", []) or [])
        methodology = payload.get("methodology", {}) or {}

        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=A4,
            leftMargin=16 * mm,
            rightMargin=16 * mm,
            topMargin=14 * mm,
            bottomMargin=14 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "AiudaTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#12263A"),
            spaceAfter=10,
            alignment=TA_LEFT,
        )
        heading_style = ParagraphStyle(
            "AiudaHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#12263A"),
            spaceAfter=8,
            spaceBefore=12,
        )
        body_style = ParagraphStyle(
            "AiudaBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#12263A"),
            spaceAfter=4,
        )
        muted_style = ParagraphStyle("AiudaMuted", parent=body_style, textColor=colors.HexColor("#536474"))

        story: List[Any] = []
        story.append(Paragraph(html.escape(str(summary.get("title", "Informe documental"))), title_style))
        story.append(Paragraph(html.escape(str(summary.get("overview", ""))), muted_style))
        story.append(Spacer(1, 6))

        teacher_issues = _filter_reportable_issues(payload)
        teacher_decision = _teacher_decision_block(payload, teacher_issues)
        teacher_priorities = _build_teacher_priorities(payload, teacher_issues, max_items=3)
        teacher_categories = _build_teacher_categories(payload, teacher_issues)
        teacher_strengths = _build_teacher_strengths(payload, teacher_categories)

        story.append(Paragraph("Informe para profesorado", heading_style))
        story.append(Paragraph(f"<b>Resultado para docencia:</b> {html.escape(str(teacher_decision.get('label', '-')))}", body_style))
        story.append(Paragraph(f"<b>Decisión rápida:</b> {html.escape(str(teacher_decision.get('decision', '-')))}", body_style))
        story.append(Paragraph(html.escape(str(teacher_decision.get('summary', '-'))), body_style))
        story.append(Paragraph(f"<b>Foco principal:</b> {html.escape(str(teacher_decision.get('focus', '-')))}", body_style))

        story.append(Paragraph("Qué conviene corregir primero", heading_style))
        if teacher_priorities:
            for item in teacher_priorities:
                story.append(Paragraph(f"<b>{html.escape(str(item.get('location', '-')))}</b>", body_style))
                story.append(Paragraph(f"<b>Área:</b> {html.escape(str(item.get('area', '-')))}", muted_style))
                story.append(Paragraph(f"<b>Qué ocurre:</b> {html.escape(str(item.get('what', '-')))}", body_style))
                story.append(Paragraph(f"<b>Qué hacer:</b> {html.escape(str(item.get('action', '-')))}", body_style))
                preview_path = item.get("preview_image_path")
                if preview_path:
                    abs_preview = pdf_path.parent / str(preview_path)
                    if abs_preview.exists():
                        try:
                            img = RLImage(str(abs_preview))
                            max_w = 72 * mm
                            max_h = 48 * mm
                            scale = min(max_w / float(img.imageWidth or max_w), max_h / float(img.imageHeight or max_h), 1.0)
                            img.drawWidth = float(img.imageWidth or max_w) * scale
                            img.drawHeight = float(img.imageHeight or max_h) * scale
                            story.append(Spacer(1, 2))
                            story.append(img)
                            if item.get("preview_caption"):
                                story.append(Paragraph(html.escape(str(item.get("preview_caption"))), muted_style))
                        except Exception:
                            pass
                story.append(Spacer(1, 4))
        else:
            story.append(Paragraph("No se detectaron incidencias prioritarias en este análisis automático.", body_style))

        story.append(Paragraph("Resumen rápido por categorías", heading_style))
        for category in teacher_categories:
            story.append(Paragraph(f"<b>{html.escape(str(category.get('title', '-')))}</b>: {html.escape(str(category.get('status', '-')))}", body_style))
            story.append(Paragraph(html.escape(str(category.get('detail', '-'))), muted_style))

        story.append(Paragraph("Qué ya funciona bien", heading_style))
        if teacher_strengths:
            for item in teacher_strengths:
                story.append(Paragraph(f"• {html.escape(str(item))}", body_style))
        else:
            story.append(Paragraph("El material presenta una base general utilizable, aunque conviene revisar algunos detalles antes de clase.", muted_style))

        story.append(Paragraph("Recomendaciones prácticas", heading_style))
        for item in recommendations[:5]:
            story.append(Paragraph(f"• {html.escape(str(item))}", body_style))
        story.append(Paragraph("A continuación se incluye un anexo técnico con el detalle completo del análisis automático, para consulta y revisión más especializada.", muted_style))
        story.append(Spacer(1, 8))

        totals = summary.get("totals", {}) or {}
        severity_totals = summary.get("severity_totals", {}) or {}
        story.append(Paragraph("Anexo técnico", heading_style))
        summary_table = Table(
            [
                ["Unidades", totals.get("units_analyzed", "-")],
                ["Alertas visuales", totals.get("reportable_issue_count", totals.get("issue_count", "-"))],
                ["Notas de procesamiento", totals.get("processing_note_count", 0)],
                ["Bloques textuales", totals.get("text_blocks_analyzed", "-")],
                ["Figuras analizadas", totals.get("figures_analyzed", "-")],
                ["Severidad alta/media/baja", f"{severity_totals.get('alta', 0)}/{severity_totals.get('media', 0)}/{severity_totals.get('baja', 0)}"],
            ],
            colWidths=[54 * mm, 118 * mm],
        )
        summary_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#DDE6EE")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDE6EE")),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#12263A")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(Paragraph("Resumen ejecutivo", heading_style))
        story.append(summary_table)

        story.append(Paragraph("Incidencias prioritarias", heading_style))
        if issues:
            for issue in issues:
                lines = [
                    Paragraph(f"<b>{html.escape(str(issue.get('location', '-')))}</b>", body_style),
                    Paragraph(html.escape(" · ".join(str(x) for x in [issue.get("category", "-"), issue.get("severity", "-")] if x)), muted_style),
                    Paragraph(html.escape(str(issue.get("message", "-"))), body_style),
                    Paragraph(f"<b>Recomendación:</b> {html.escape(str(issue.get('recommendation', '-')))}", body_style),
                ]
                metric_parts = []
                if issue.get("ratio") is not None and issue.get("threshold") is not None:
                    metric_parts.append(f"Ratio {issue.get('ratio')}:1")
                    metric_parts.append(f"Umbral {issue.get('threshold')}:1")
                if issue.get("score_ratio") is not None:
                    metric_parts.append(f"Score CVD {issue.get('score_ratio')}")
                if issue.get("bbox_human"):
                    metric_parts.append(f"Zona {issue.get('bbox_human')}")
                if metric_parts:
                    lines.append(Paragraph(html.escape(" · ".join(metric_parts)), muted_style))
                if issue.get("excerpt"):
                    lines.append(Paragraph(f"<b>Extracto:</b> {html.escape(str(issue.get('excerpt')))}", muted_style))
                if issue.get("reference_code"):
                    lines.append(Paragraph(f"<b>Referencia:</b> {html.escape(str(issue.get('reference_code')))}", muted_style))
                issue_table = Table([[lines]], colWidths=[178 * mm])
                issue_table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#DDE6EE")),
                            ("PADDING", (0, 0), (-1, -1), 10),
                        ]
                    )
                )
                story.append(issue_table)
                story.append(Spacer(1, 4))
        else:
            story.append(Paragraph("No se detectaron incidencias visuales prioritarias en este análisis automático.", body_style))

        story.append(Paragraph("Recomendaciones generales", heading_style))
        for item in recommendations:
            story.append(Paragraph(f"• {html.escape(str(item))}", body_style))

        story.append(Paragraph("Métricas visuales", heading_style))
        text_compact_note = build_text_metric_compact_note(visual_metrics)
        figure_compact_note = build_figure_metric_compact_note(visual_metrics)

        if text_compact_note:
            story.append(Paragraph(f"• Contraste textual: {html.escape(text_compact_note)}", body_style))
        else:
            for line in build_metric_summary_lines(visual_metrics.get("text_metrics_summary", {}) or {}, "text"):
                story.append(Paragraph(f"• Contraste textual: {html.escape(line)}", body_style))
            text_rows = _select_relevant_metric_rows(visual_metrics.get("text_metric_samples", []) or [])
            if text_rows:
                story.append(Paragraph("Muestras de contraste textual", heading_style))
                table = Table(
                    [["Ubicación", "Bloque", "Estado", "Ratio", "Umbral", "Zona"]] + [
                        [
                            str(row.get("location", "-")),
                            str(row.get("text_block_index", "-")),
                            str(row.get("status", "-")),
                            str(row.get("ratio", "-")),
                            str(row.get("threshold", "-")),
                            str(row.get("bbox_human", "-")),
                        ]
                        for row in text_rows
                    ],
                    colWidths=[48 * mm, 14 * mm, 28 * mm, 18 * mm, 18 * mm, 46 * mm],
                    repeatRows=1,
                )
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE6EE")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDE6EE")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(table)

        if figure_compact_note:
            story.append(Paragraph(f"• Figuras/CVD: {html.escape(figure_compact_note)}", body_style))
        else:
            for line in build_metric_summary_lines(visual_metrics.get("figure_metrics_summary", {}) or {}, "figure"):
                story.append(Paragraph(f"• Figuras/CVD: {html.escape(line)}", body_style))
            figure_rows = _select_relevant_metric_rows(visual_metrics.get("figure_metric_samples", []) or [])
            if figure_rows:
                story.append(Paragraph("Muestras de figuras y CVD", heading_style))
                table = Table(
                    [["Ubicación", "Figura", "Estado", "Ratio", "Peor CVD", "Zona"]] + [
                        [
                            str(row.get("location", "-")),
                            str(row.get("figure_index", "-")),
                            str(row.get("status", "-")),
                            str(row.get("ratio", "-")),
                            str(row.get("worst_cvd_ratio", "-")),
                            str(row.get("bbox_human", "-")),
                        ]
                        for row in figure_rows
                    ],
                    colWidths=[52 * mm, 14 * mm, 26 * mm, 18 * mm, 22 * mm, 40 * mm],
                    repeatRows=1,
                )
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE6EE")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDE6EE")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(table)


        processing_notes = clean_list_items(visual_metrics.get("processing_notes", []) or [])
        if processing_notes:
            story.append(Paragraph("Notas de procesamiento", heading_style))
            for note in processing_notes:
                story.append(Paragraph(f"• {html.escape(str(note))}", muted_style))

        story.append(Paragraph(methodology.get("title", "Metodología y referencias"), heading_style))
        for rule in methodology.get("contrast_rules", []) or []:
            rule_label = html.escape(str(rule.get("label", "-")))
            detail = html.escape(str(rule.get("details", "")))
            ref = html.escape(str(rule.get("reference_code", "-")))
            threshold = rule.get("threshold")
            threshold_text = f" Umbral: {threshold}:1." if threshold is not None else ""
            story.append(Paragraph(f"<b>{rule_label}</b> · {html.escape(str(rule.get('metric_type', '-')))} · {ref}.{threshold_text}", body_style))
            if detail:
                story.append(Paragraph(detail, muted_style))
        limitations = clean_list_items(methodology.get("limitations", []) or [])
        if limitations:
            story.append(Paragraph("Limitaciones de la estimación", heading_style))
            for item in limitations:
                story.append(Paragraph(f"• {html.escape(str(item))}", muted_style))
        refs = methodology.get("references", []) or []
        if refs:
            story.append(Paragraph("Referencias", heading_style))
            for ref in refs:
                story.append(Paragraph(f"• {html.escape(str(ref.get('title', '-')))}", body_style))
                if ref.get("details"):
                    story.append(Paragraph(html.escape(str(ref.get("details"))), muted_style))

        doc.build(story)
        return
    except Exception as exc:
        raise RuntimeError(f"No se pudo generar el PDF del informe: {exc}")

def save_report_variants(
    payload: dict,
    stem: str,
    output_dir: Path,
    variant: str,
) -> Dict[str, str]:
    txt_name = f"{stem}.{variant}.txt"
    json_name = f"{stem}.{variant}.json"
    html_name = f"{stem}.{variant}.html"
    pdf_name = f"{stem}.{variant}.pdf"

    txt_path = output_dir / txt_name
    json_path = output_dir / json_name
    html_path = output_dir / html_name
    pdf_path = output_dir / pdf_name

    html_content = render_report_html(
        payload,
        title=payload.get("summary", {}).get("title", "Informe documental"),
    )

    write_text(txt_path, render_report_text(payload))
    write_json(json_path, payload)
    write_text(html_path, html_content)
    render_html_to_pdf(html_content, pdf_path, payload=payload)

    return {
        "txt": txt_name,
        "json": json_name,
        "html": html_name,
        "pdf": pdf_name,
    }


# ---------------------------------------------------------
# MAIN PROCESS
# ---------------------------------------------------------


def process_document(
    input_file: Path,
    output_dir: Path,
    task_id: str,
    source_lang: Optional[str],
    target_langs: List[str],
    use_ocr: bool,
    max_translation_chars: int,
) -> int:
    ensure_dir(output_dir)
    logger = build_logger(output_dir)
    stem = sanitize_stem(input_file.stem)

    write_status(
        output_dir,
        task_id=task_id,
        state="processing",
        progress=5,
        stage="startup",
        message="Inicializando tarea documental",
        extra={"input_file": str(input_file)},
    )

    logger.info("%s", APP_NAME)
    logger.info("Archivo de entrada: %s", input_file)
    logger.info("Directorio de salida: %s", output_dir)

    write_status(
        output_dir,
        task_id=task_id,
        state="processing",
        progress=20,
        stage="extraction",
        message="Extrayendo contenido del documento",
    )

    document_type, units = extract_document_units(input_file, use_ocr=use_ocr, logger=logger)
    full_text = normalize_whitespace("\n\n".join(unit.get("text", "") for unit in units))
    detected_lang = detect_language_heuristic(full_text)
    manual_source_lang = normalize_lang_code(source_lang)
    used_source_lang = manual_source_lang or detected_lang or "es"
    logger.info("Idioma detectado/usado para metadatos: %s", used_source_lang)

    write_status(
        output_dir,
        task_id=task_id,
        state="processing",
        progress=50,
        stage="analysis",
        message="Analizando accesibilidad documental",
        extra={"source_language_used": used_source_lang},
    )

    report_base = analyze_document(units, document_type=document_type, input_filename=input_file.name)
    original_payload: Dict[str, Any] = {
        "task_id": task_id,
        "input_file": str(input_file),
        "input_filename": input_file.name,
        "document_type": document_type,
        "source_language_requested": manual_source_lang,
        "source_language_detected": detected_lang,
        "source_language_used": used_source_lang,
        "report_language": "es",
        **report_base,
    }

    original_payload["teacher_priority_assets"] = _build_teacher_priority_assets(
        input_file=input_file,
        output_dir=output_dir,
        payload=original_payload,
        logger=logger,
        max_items=3,
    )

    original_outputs = save_report_variants(original_payload, stem=stem, output_dir=output_dir, variant="original")
    logger.info("Informe original generado.")

    clean_targets: List[str] = []
    for lang in target_langs:
        lang = normalize_lang_code(lang)
        if lang in ALLOWED_TARGETS and lang not in clean_targets:
            clean_targets.append(lang)

    translated_outputs: Dict[str, Dict[str, str]] = {}
    translator: Optional[NllbTranslator] = None

    if clean_targets:
        write_status(
            output_dir,
            task_id=task_id,
            state="processing",
            progress=75,
            stage="translation",
            message="Generando informes multilingües",
            extra={"source_language_used": used_source_lang},
        )

        need_translation = any(lang != "es" for lang in clean_targets)
        if need_translation:
            try:
                translator = NllbTranslator(model_name=DEFAULT_TRANSLATION_MODEL)
            except Exception as exc:
                logger.warning("No se pudo inicializar NLLB. Se crearán copias no traducidas de los informes: %s", exc)
                translator = None

        for idx, tgt_lang in enumerate(clean_targets, start=1):
            logger.info("Generando informe %s/%s: %s", idx, len(clean_targets), tgt_lang)
            if tgt_lang == "es":
                translated_payload = deepcopy(original_payload)
            else:
                translated_payload = translate_payload_strings(
                    deepcopy(original_payload),
                    translator=translator,
                    tgt_lang=tgt_lang,
                    logger=logger,
                    max_chars=max_translation_chars,
                )
            translated_payload["report_language"] = tgt_lang
            translated_payload["target_language"] = tgt_lang
            translated_outputs[tgt_lang] = save_report_variants(
                translated_payload,
                stem=stem,
                output_dir=output_dir,
                variant=tgt_lang,
            )

    success_message = "Procesamiento documental completado correctamente."
    result_payload = {
        "ok": True,
        "task_id": task_id,
        "status": "finished",
        "message": success_message,
        "error": None,
        "input_file": str(input_file),
        "document_type": document_type,
        "source_language_requested": manual_source_lang,
        "source_language_detected": detected_lang,
        "source_language_used": used_source_lang,
        "target_languages_requested": clean_targets,
        "outputs": collect_output_filenames(output_dir),
        "original": original_outputs,
        "translations": translated_outputs,
    }

    write_json(output_dir / "result.json", result_payload)

    write_status(
        output_dir,
        task_id=task_id,
        state="finished",
        progress=100,
        stage="done",
        message=success_message,
        extra={
            "ok": True,
            "error": None,
            "source_language_used": used_source_lang,
            "result_file": "result.json",
        },
    )

    logger.info("Proceso documental completado correctamente.")
    return 0



def main() -> int:
    args = parse_args()

    input_file = Path(args.input_file).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    if not input_file.exists():
        print(f"[ERROR] No existe el archivo: {input_file}", file=sys.stderr)
        return 1
    if not input_file.is_file():
        print(f"[ERROR] La entrada no es un archivo válido: {input_file}", file=sys.stderr)
        return 1

    try:
        return process_document(
            input_file=input_file,
            output_dir=output_dir,
            task_id=args.task_id,
            source_lang=args.source_lang,
            target_langs=args.translate_to,
            use_ocr=bool(args.ocr),
            max_translation_chars=args.max_translation_chars,
        )
    except Exception as exc:
        ensure_dir(output_dir)
        error_message = f"Error al procesar el documento: {exc}"
        write_json(
            output_dir / "result.json",
            {
                "ok": False,
                "task_id": args.task_id,
                "status": "error",
                "message": error_message,
                "error": str(exc),
                "outputs": collect_output_filenames(output_dir),
            },
        )
        write_status(
            output_dir,
            task_id=args.task_id,
            state="error",
            progress=100,
            stage="failed",
            message=error_message,
            extra={
                "ok": False,
                "error": str(exc),
                "result_file": "result.json",
            },
        )
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
